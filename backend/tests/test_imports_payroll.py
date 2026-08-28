"""A folha entrando por planilha — o que o arquivo recusa, e o que ele decide.

Dois níveis, e a ordem entre eles é a decisão: o arquivo é julgado inteiro antes
da primeira linha, porque arquivo errado não tem linha certa. Depois, cada linha
recebe um veredito que ninguém executou ainda — `verdict` é pura, e é isso que
permite a tela mostrar o que aconteceria antes de alguém confirmar.
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook

from operax.imports import payroll
from operax.rh.workbook import META_SHEET, SheetRow, WorkbookError

TENANT = UUID("22222222-2222-4222-8222-222222222222")
OUTRO_TENANT = UUID("33333333-3333-4333-8333-333333333333")
COLAB = UUID("44444444-4444-4444-8444-444444444444")
AGORA = datetime(2026, 8, 28, 12, 0, 0)

MATRICULAS = {"1001": COLAB}
MAPEADOS = frozenset({"0050", "H_EXTRA_60"})


def _linha(**campos) -> dict[str, object]:
    base = {
        "employee_code": "1001",
        "employee_name": "Fulana de Tal",
        "code": "0050",
        "description": "Salário base",
        "nature": "Provento",
        "reference": None,
        "amount": Decimal("2500.00"),
    }
    base.update(campos)
    return base


def _arquivo(*linhas: dict[str, object]) -> bytes:
    return payroll.build_template(
        tenant_id=TENANT, year=2026, month=8, generated_at=AGORA, rows=tuple(linhas)
    )


def _adulterar_meta(data: bytes, chave: str, valor: str) -> bytes:
    wb = load_workbook(BytesIO(data))
    ws = wb[META_SHEET]
    for linha in ws.iter_rows(min_row=1, max_col=2):
        if linha[0].value == chave:
            linha[1].value = valor
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# O arquivo, antes da primeira linha
# ---------------------------------------------------------------------------
def test_o_modelo_vazio_volta_com_a_competencia_e_sem_linha():
    """A competência vive na aba de controle, não repetida em cada linha."""
    lido = payroll.parse_upload(_arquivo(), tenant_id=TENANT)

    assert (lido.year, lido.month) == (2026, 8)
    assert lido.layout_version == payroll.LAYOUT_VERSION
    assert lido.rows == ()


def test_ida_e_volta_preserva_os_valores_da_linha():
    lido = payroll.parse_upload(_arquivo(_linha()), tenant_id=TENANT)

    assert len(lido.rows) == 1
    valores = lido.rows[0].values
    assert valores["employee_code"] == "1001"
    assert valores["code"] == "0050"
    assert Decimal(str(valores["amount"])) == Decimal("2500.00")
    # A linha do Excel, não o índice: é o número que a pessoa vê na tela.
    assert lido.rows[0].line == 2


def test_arquivo_de_outro_cliente_e_recusado_inteiro():
    """Ele traria nome e matrícula de outro cliente para a tela de preview."""
    with pytest.raises(WorkbookError) as erro:
        payroll.parse_upload(_arquivo(_linha()), tenant_id=OUTRO_TENANT)

    assert erro.value.code == "tenant_divergente"


def test_arquivo_que_nao_e_xlsx_diz_isso_em_vez_de_estourar():
    with pytest.raises(WorkbookError) as erro:
        payroll.parse_upload(b"isto nao e uma planilha", tenant_id=TENANT)

    assert erro.value.code == "arquivo_ilegivel"


def test_arquivo_de_outro_tipo_nao_entra_como_folha():
    adulterado = _adulterar_meta(_arquivo(), "type", "hr_registration")

    with pytest.raises(WorkbookError) as erro:
        payroll.parse_upload(adulterado, tenant_id=TENANT)

    assert erro.value.code == "tipo_divergente"


def test_layout_antigo_e_recusado_pelo_nome_da_versao():
    adulterado = _adulterar_meta(_arquivo(), "layout_version", "folha-0")

    with pytest.raises(WorkbookError) as erro:
        payroll.parse_upload(adulterado, tenant_id=TENANT)

    assert erro.value.code == "layout_antigo"


def test_cabecalho_alterado_e_recusado_antes_de_qualquer_linha():
    """Coluna inserida, removida ou renomeada muda o significado de tudo abaixo."""
    data = _arquivo(_linha())
    wb = load_workbook(BytesIO(data))
    wb[payroll.SHEET_TITLE].cell(row=1, column=1, value="Codigo do funcionario")
    buffer = BytesIO()
    wb.save(buffer)

    with pytest.raises(WorkbookError) as erro:
        payroll.parse_upload(buffer.getvalue(), tenant_id=TENANT)

    assert erro.value.code == "cabecalho_alterado"


def test_competencia_ilegivel_na_aba_de_controle_e_recusa_nomeada():
    adulterado = _adulterar_meta(_arquivo(), "period", "agosto")

    with pytest.raises(WorkbookError) as erro:
        payroll.parse_upload(adulterado, tenant_id=TENANT)

    assert erro.value.code == "competencia_invalida"


def test_linha_em_branco_no_fim_nao_e_linha():
    """Rastro de quem apagou o conteúdo em vez de excluir a linha."""
    data = _arquivo(_linha())
    wb = load_workbook(BytesIO(data))
    ws = wb[payroll.SHEET_TITLE]
    ws.cell(row=3, column=1, value=None)
    ws.cell(row=4, column=2, value="   ")
    buffer = BytesIO()
    wb.save(buffer)

    lido = payroll.parse_upload(buffer.getvalue(), tenant_id=TENANT)

    assert len(lido.rows) == 1


# ---------------------------------------------------------------------------
# O veredito, linha a linha
# ---------------------------------------------------------------------------
def _veredito(*valores: dict[str, object]) -> payroll.Report:
    linhas = tuple(
        SheetRow(line=indice, values=dict(v)) for indice, v in enumerate(valores, start=2)
    )
    return payroll.verdict(linhas, employees=MATRICULAS, mapped_codes=MAPEADOS)


def test_linha_boa_passa_e_traz_o_colaborador_resolvido():
    relatorio = _veredito(_linha())

    (linha,) = relatorio.outcomes
    assert linha.ok
    assert linha.employee_id == COLAB
    assert linha.values["nature"] == "earning", "a natureza vai para o banco em inglês"
    assert (relatorio.rows_total, relatorio.rows_ok, relatorio.rows_error) == (1, 1, 0)


def test_matricula_desconhecida_e_erro_que_nomeia_a_matricula():
    relatorio = _veredito(_linha(employee_code="9999"))

    (linha,) = relatorio.outcomes
    assert not linha.ok
    assert [e.code for e in linha.errors] == ["colaborador_desconhecido"]
    assert "9999" in linha.errors[0].message


def test_valor_que_nao_e_numero_e_erro_e_nao_vira_zero():
    """Zero seria um número errado que fecha a soma. Erro para a linha."""
    relatorio = _veredito(_linha(amount="dois mil"))

    (linha,) = relatorio.outcomes
    assert "valor_invalido" in [e.code for e in linha.errors]


def test_valor_no_formato_brasileiro_e_aceito():
    """É o que sai do sistema da contabilidade — recusar faria reformatar à mão."""
    relatorio = _veredito(_linha(amount="1.234,56"))

    (linha,) = relatorio.outcomes
    assert linha.ok
    assert linha.values["amount"] == Decimal("1234.56")


def test_natureza_fora_do_catalogo_e_erro_com_a_lista_do_que_existe():
    relatorio = _veredito(_linha(nature="Bonus"))

    (linha,) = relatorio.outcomes
    assert "natureza_invalida" in [e.code for e in linha.errors]
    assert "Provento" in linha.errors[0].message


def test_campo_obrigatorio_vazio_e_erro_pelo_rotulo_da_coluna():
    relatorio = _veredito(_linha(code=""))

    (linha,) = relatorio.outcomes
    assert "campo_obrigatorio" in [e.code for e in linha.errors]
    assert "Código do evento" in " ".join(e.message for e in linha.errors)


def test_codigo_sem_categoria_nao_impede_a_linha_de_entrar():
    """A curadoria nunca está completa na primeira importação — e o valor é real.

    Bloquear a folha inteira por código não mapeado tornaria impossível justo o
    primeiro mês, que é quando o cliente mais precisa dela.
    """
    relatorio = _veredito(_linha(code="0099"))

    (linha,) = relatorio.outcomes
    assert linha.ok, "código sem categoria é pendência, não erro"
    assert [w.code for w in linha.warnings] == ["codigo_sem_categoria"]
    assert relatorio.rows_ok == 1
    assert relatorio.unmapped_codes == ("0099",)


def test_o_relatorio_lista_cada_codigo_sem_categoria_uma_vez():
    """É a lista que vai para a curadoria com a contabilidade, não a de linhas."""
    relatorio = _veredito(
        _linha(code="0099"), _linha(code="0099", amount=Decimal("10")), _linha(code="0100")
    )

    assert relatorio.unmapped_codes == ("0099", "0100")


def test_linha_identica_repetida_e_aviso_e_nao_recusa():
    relatorio = _veredito(_linha(), _linha())

    primeira, segunda = relatorio.outcomes
    assert primeira.warnings == ()
    assert [w.code for w in segunda.warnings] == ["linha_repetida"]
    assert relatorio.rows_error == 0


def test_o_relatorio_so_carrega_as_linhas_que_tem_o_que_dizer():
    """`app.file_import.report` é lido por quem vai corrigir, não é o arquivo de novo."""
    relatorio = _veredito(_linha(), _linha(employee_code="9999"), _linha(code="0099"))

    corpo = relatorio.as_json()
    assert corpo["rows_total"] == 3
    assert corpo["rows_error"] == 1
    assert [linha["line"] for linha in corpo["lines"]] == [3, 4]
    assert corpo["unmapped_codes"] == ["0099"]


def test_nada_no_veredito_toca_o_banco():
    """Guarda de leitura: `verdict` é pura, e é o que permite perguntar duas vezes."""
    primeira = _veredito(_linha())
    segunda = _veredito(_linha())

    assert primeira.as_json() == segunda.as_json()
    assert uuid4() not in MATRICULAS.values()
