"""O gate do R2: modelo → edição → upload → preview → confirmação parcial → reimport.

Nada aqui abre banco, mas nada aqui finge SQL: `tenant_scope` e `user_scope` são
as duas costuras, e o que entra no lugar delas é um cursor de mentira que executa
as instruções de verdade — inclusive `bind_tenant`, que recusa qualquer statement
sem filtro de tenant. O que as policies devolvem é provado em SQL, na suíte de
isolamento; o que este módulo faz com a resposta é provado aqui.

O caminho inteiro é um teste só, e de propósito: o valor do R2 não está em
nenhuma das etapas isoladas, e sim em o arquivo voltar, três linhas caírem, as
outras entrarem, e o reenvio das corrigidas **não** reescrever as que já estavam
certas.
"""

from __future__ import annotations

from typing import Any
from uuid import UUID, uuid4

import pytest
from fastapi.testclient import TestClient

from operax.core.tenant import bind_tenant
from operax.rh import repository
from operax.rh.workbook import CONTENT_TYPE, META_SHEET
from server.main import app
from server.routers import rh

PESSOAS = [
    {"employee_id": uuid4(), "registration_number": f"100{n}", "name": nome, "hr_code": None}
    for n, nome in enumerate(
        [
            "Ana Personagem",
            "Bruno Personagem",
            "Carla Personagem",
            "Davi Personagem",
            "Elis Personagem",
        ],
        start=1,
    )
]


class FakeCursor:
    """Um banco de mentira endereçado por statement, não por ordem de chamada.

    Endereçar por ordem tornaria o teste refém do encadeamento interno do
    endpoint: mover uma leitura de lugar quebraria o teste sem quebrar o produto.
    """

    def __init__(self, estado: FakeDB, context: Any, checar_tenant: bool) -> None:
        self._estado = estado
        self._context = context
        self._checar_tenant = checar_tenant
        self._resultado: Any = None

    async def execute(self, statement: str, params: Any = None) -> None:
        if self._checar_tenant:
            # A mesma recusa do `TenantScope`: statement sem filtro de tenant não
            # roda. Se um dia alguém escrever um update sem ele, é aqui que para.
            params = bind_tenant(statement, params, self._context)
        params = params or {}
        self._estado.statements.append((" ".join(statement.split()), dict(params)))
        self._resultado = self._estado.responder(statement, params)

    async def fetchone(self) -> Any:
        return self._resultado[0] if self._resultado else None

    async def fetchall(self) -> list[dict[str, Any]]:
        return list(self._resultado or [])


class FakeScope:
    def __init__(self, cursor: FakeCursor) -> None:
        self._cursor = cursor

    async def __aenter__(self) -> FakeCursor:
        return self._cursor

    async def __aexit__(self, *exc: object) -> None:
        return None


class FakeDB:
    """Estado do cliente: pessoas, importações e o log de auditoria."""

    def __init__(self, *, admin: bool = True, dominios: bool = True) -> None:
        self.pessoas = [dict(p) for p in PESSOAS]
        self.imports: dict[UUID, dict[str, Any]] = {}
        self.audit: list[dict[str, Any]] = []
        self.statements: list[tuple[str, dict[str, Any]]] = []
        self.admin = admin
        self.dominios = dominios

    # -- leitura ------------------------------------------------------------
    def responder(self, statement: str, params: dict[str, Any]) -> Any:
        sql = " ".join(statement.split())
        if "util.is_admin" in sql:
            return [
                {
                    "admin": self.admin,
                    "pii": self.dominios,
                    "compensation": self.dominios,
                    "health": self.dominios,
                }
            ]
        if "from app.payroll_period" in sql:
            return [{"ends_on": None}]
        if "from app.employee e" in sql:
            return [dict(p) for p in self.pessoas]
        if "from app.file_import" in sql:
            return [self.imports.get(params["import_id"])] if params.get("import_id") else []
        if "insert into app.file_import" in sql:
            self.imports[params["import_id"]] = {
                "id": params["import_id"],
                "type": params["type"],
                "storage_path": params["storage_path"],
                "file_name": params["file_name"],
                "layout_version": params["layout_version"],
                "status": "received",
                "rows_total": None,
                "rows_ok": None,
                "rows_error": None,
            }
            return []
        if "update app.file_import" in sql:
            registro = self.imports[params["import_id"]]
            registro |= {
                "status": params["status"],
                "rows_total": params["rows_total"],
                "rows_ok": params["rows_ok"],
                "rows_error": params["rows_error"],
            }
            return []
        if "insert into app.audit_log" in sql:
            self.audit.append(dict(params))
            return []
        if "update app.employee set" in sql or "insert into app.employee_pii" in sql:
            self._gravar(params)
            return [{"id": params["employee_id"]}]
        if "app.employee_compensation" in sql:
            self._gravar(params)
            return [{"id": uuid4()}]
        return []

    def _gravar(self, params: dict[str, Any]) -> None:
        for pessoa in self.pessoas:
            if pessoa["employee_id"] == params["employee_id"]:
                for chave, valor in params.items():
                    if chave not in {"employee_id", "tenant_id", "user_id"}:
                        pessoa[chave] = valor


class FakeStore:
    """Storage em memória. O confirm relê daqui, como releria do Supabase."""

    def __init__(self) -> None:
        self.objetos: dict[str, bytes] = {}

    async def put(self, path: str, data: bytes, content_type: str) -> None:
        self.objetos[path] = data

    async def get(self, path: str) -> bytes:
        return self.objetos[path]


@pytest.fixture
def db(monkeypatch: pytest.MonkeyPatch):
    def instalar(**kwargs: Any) -> FakeDB:
        estado = FakeDB(**kwargs)
        monkeypatch.setattr(
            repository,
            "tenant_scope",
            lambda context, schema="app": FakeScope(FakeCursor(estado, context, True)),
        )
        monkeypatch.setattr(
            repository,
            "user_scope",
            lambda context, schema="app": FakeScope(FakeCursor(estado, context, False)),
        )
        return estado

    return instalar


@pytest.fixture
def store() -> FakeStore:
    memoria = FakeStore()
    app.dependency_overrides[rh.get_store] = lambda: memoria
    return memoria


def auth(issue_token) -> dict[str, str]:
    return {"Authorization": f"Bearer {issue_token()}"}


def baixar(client: TestClient, issue_token, tipo: str = "hr_link"):
    return client.get(f"/rh/template/{tipo}", headers=auth(issue_token))


def editar(conteudo: bytes, edicoes: dict[int, dict[str, Any]]) -> bytes:
    """Preenche células como o usuário faria — pelo rótulo da coluna."""
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(conteudo))
    ws = next(wb[nome] for nome in wb.sheetnames if nome != META_SHEET)
    rotulos = {celula.value: celula.column for celula in next(ws.iter_rows(min_row=1, max_row=1))}
    for linha, valores in edicoes.items():
        for rotulo, valor in valores.items():
            ws.cell(row=linha, column=rotulos[rotulo]).value = valor
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def enviar(client: TestClient, issue_token, conteudo: bytes, tipo: str = "hr_link"):
    return client.post(
        "/rh/imports",
        headers=auth(issue_token),
        data={"tipo": tipo},
        files={"arquivo": ("modelo.xlsx", conteudo, CONTENT_TYPE)},
    )


# ---------------------------------------------------------------------------
# O caminho inteiro
# ---------------------------------------------------------------------------
def test_do_modelo_ate_a_confirmacao_parcial_e_o_reimport(
    client: TestClient, issue_token, db, store
):
    estado = db()

    modelo = baixar(client, issue_token)
    assert modelo.status_code == 200
    assert modelo.headers["content-type"] == CONTENT_TYPE
    assert "hr_link.xlsx" in modelo.headers["content-disposition"]
    # O download em si é registrado: o arquivo leva o dado embora.
    assert [a["action"] for a in estado.audit] == ["export"]

    # Duas linhas certas e três erradas, cada uma de um jeito que a planilha real
    # produz: ID RH repetido, nome do ponto reescrito, matrícula colada errada.
    preenchido = editar(
        modelo.content,
        {
            2: {"ID RH": "RH-01"},
            3: {"ID RH": "RH-02"},
            4: {"ID RH": "RH-01"},
            5: {"ID RH": "RH-04", "Nome": "Davi P. Personagem"},
            6: {"ID RH": "RH-05", "Matrícula": "9999"},
        },
    )

    preview = enviar(client, issue_token, preenchido)
    assert preview.status_code == 201
    corpo = preview.json()
    assert corpo["counts"] == {"total": 5, "ok": 2, "unchanged": 0, "error": 3}
    assert [linha["status"] for linha in corpo["lines"]] == [
        "ok",
        "ok",
        "error",
        "error",
        "error",
    ]
    assert [linha["errors"][0]["code"] for linha in corpo["lines"] if linha["errors"]] == [
        "duplicado_no_arquivo",
        "campo_do_sync",
        "chave_desconhecida",
    ]
    # Preview não grava: ninguém ganhou ID RH e nada foi auditado além do export.
    assert all(p["hr_code"] is None for p in estado.pessoas)
    assert [a["action"] for a in estado.audit] == ["export"]

    import_id = corpo["import_id"]
    confirmado = client.post(f"/rh/imports/{import_id}/confirm", headers=auth(issue_token))
    assert confirmado.status_code == 200
    resultado = confirmado.json()
    assert resultado["applied"] == 2
    assert resultado["partial"] is True
    assert resultado["status"] == "processed"

    assert [p["hr_code"] for p in estado.pessoas] == ["RH-01", "RH-02", None, None, None]
    gravacoes = [a for a in estado.audit if a["action"] == "update"]
    assert len(gravacoes) == 2
    for registro in gravacoes:
        assert registro["depois"].obj["_origem"]["file_import_id"] == import_id

    # `processed` com `rows_error > 0` é o "parcial" da SPEC §3 — não há sexto
    # valor no check de status.
    guardado = estado.imports[UUID(import_id)]
    assert (guardado["status"], guardado["rows_ok"], guardado["rows_error"]) == ("processed", 2, 3)

    # --- o reimport das corrigidas ----------------------------------------
    segundo_modelo = baixar(client, issue_token)
    corrigido = editar(
        segundo_modelo.content,
        {4: {"ID RH": "RH-03"}, 5: {"ID RH": "RH-04"}, 6: {"ID RH": "RH-05"}},
    )

    segundo_preview = enviar(client, issue_token, corrigido)
    assert segundo_preview.json()["counts"] == {
        "total": 5,
        "ok": 3,
        # As duas que já entraram voltam como "não mudou": o reenvio não as
        # reescreve, e é isso que faz o reimport ser seguro.
        "unchanged": 2,
        "error": 0,
    }

    segundo_id = segundo_preview.json()["import_id"]
    fechamento = client.post(f"/rh/imports/{segundo_id}/confirm", headers=auth(issue_token))
    assert fechamento.json()["applied"] == 3
    assert fechamento.json()["partial"] is False
    assert [p["hr_code"] for p in estado.pessoas] == [
        "RH-01",
        "RH-02",
        "RH-03",
        "RH-04",
        "RH-05",
    ]


# ---------------------------------------------------------------------------
# Recusas
# ---------------------------------------------------------------------------
def test_confirmar_duas_vezes_nao_grava_de_novo(client: TestClient, issue_token, db, store):
    estado = db()
    preenchido = editar(baixar(client, issue_token).content, {2: {"ID RH": "RH-01"}})
    import_id = enviar(client, issue_token, preenchido).json()["import_id"]

    primeiro = client.post(f"/rh/imports/{import_id}/confirm", headers=auth(issue_token))
    assert primeiro.status_code == 200
    repetido = client.post(f"/rh/imports/{import_id}/confirm", headers=auth(issue_token))

    assert repetido.status_code == 409
    assert len([a for a in estado.audit if a["action"] == "update"]) == 1


def test_arquivo_que_nao_saiu_do_sistema_e_recusado_e_fica_guardado(
    client: TestClient, issue_token, db, store
):
    estado = db()
    resposta = enviar(client, issue_token, b"matricula;nome;id_rh")

    assert resposta.status_code == 422
    assert "modelo atual" in resposta.json()["detail"].lower()
    # O arquivo recusado continua no storage e o registro fica com o motivo: é a
    # prova do que foi enviado.
    assert len(store.objetos) == 1
    assert [i["status"] for i in estado.imports.values()] == ["validation_error"]


def test_quem_nao_e_do_dp_nao_baixa_modelo(client: TestClient, issue_token, db, store):
    db(admin=False)

    assert baixar(client, issue_token).status_code == 403


def test_sem_o_dominio_de_remuneracao_o_modelo_de_salario_nao_desce(
    client: TestClient, issue_token, db, store
):
    db(dominios=False)

    assert baixar(client, issue_token, "hr_compensation").status_code == 403
    # E o de vínculo, que não carrega domínio sensível, continua descendo.
    assert baixar(client, issue_token, "hr_link").status_code == 200


def test_tipo_sem_caminho_de_volta_responde_501_com_o_motivo(
    client: TestClient, issue_token, db, store
):
    db()
    resposta = baixar(client, issue_token, "hr_leave")

    assert resposta.status_code == 501
    assert "policy" in resposta.json()["detail"]


def test_tipo_inexistente_e_404(client: TestClient, issue_token, db, store):
    db()

    assert baixar(client, issue_token, "hr_qualquer").status_code == 404
