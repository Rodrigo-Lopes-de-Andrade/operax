"""A planilha que sai e a planilha que volta — e o que ela recusa.

O censo é a verdade de referência do G4 desde 09/09/2026, então um erro aqui não
produz uma tela feia: produz uma taxa de falso positivo errada, e é ela que
autoriza (ou não) o primeiro alerta a chegar num gestor. Por isso metade destes
testes é sobre recusar arquivo, e cada recusa vem com a linha equivalente que
precisa continuar passando.

A outra metade é sobre a cobertura. Uma taxa sobre censo parcial é o falso verde
que este repositório vive achando, e `Measurement.gate_passes` devolve `None`
para dizer "ainda não sei" — nunca `True` por silêncio.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook

from operax.motor.adjudicacao import (
    Measurement,
    Slice,
    build_workbook,
    parse,
)

EVENTO_A = UUID("aaaaaaaa-0000-0000-0000-00000000000a")
EVENTO_B = UUID("bbbbbbbb-0000-0000-0000-00000000000b")


def linha(evento: UUID, **overrides: Any) -> dict[str, Any]:
    base: dict[str, Any] = {
        "deviation_event_id": evento,
        "employee_name": "Ana Personagem",
        "unit_name": "A Centro",
        "reference_date": date(2026, 8, 31),
        "type": "late_entry",
        "minutes": -15,
        "expected_time": "08:00",
        "actual_time": "08:15",
        "punches": "08:15 12:00 13:00 18:00",
        "workload_minutes": 480,
        "confidence": 100,
        "verdict": None,
        "cause": None,
        "note": None,
    }
    return {**base, **overrides}


def preencher(data: bytes, respostas: dict[int, tuple[str, str, str]]) -> bytes:
    """Faz o que o julgador faz: escreve veredito, causa e observação."""
    wb = load_workbook(BytesIO(data))
    ws = wb[wb.sheetnames[0]]
    for line, (veredito, causa, nota) in respostas.items():
        ws.cell(row=line, column=12, value=veredito)
        ws.cell(row=line, column=13, value=causa)
        ws.cell(row=line, column=14, value=nota)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Ida e volta
# ---------------------------------------------------------------------------


def test_ida_e_volta_devolve_o_veredito_traduzido() -> None:
    planilha = build_workbook([linha(EVENTO_A), linha(EVENTO_B)])
    preenchida = preencher(
        planilha,
        {
            2: ("verdadeiro", "", "atraso real"),
            3: ("falso", "não bate ponto", "supervisor"),
        },
    )

    linhas, erros = parse(preenchida)

    assert erros == []
    assert [(linha_.deviation_event_id, linha_.verdict, linha_.cause) for linha_ in linhas] == [
        (EVENTO_A, "true_positive", None),
        (EVENTO_B, "false_positive", "exempt_from_punching"),
    ]
    assert linhas[1].note == "supervisor"


def test_o_ja_julgado_volta_preenchido_na_planilha() -> None:
    """Reexportar no meio do censo não pode apagar o trabalho de quem já julgou."""
    planilha = build_workbook(
        [linha(EVENTO_A, verdict="false_positive", cause="wrong_schedule", note="12x36")]
    )
    ws = load_workbook(BytesIO(planilha))["censo"]

    assert ws.cell(row=2, column=12).value == "falso"
    assert ws.cell(row=2, column=13).value == "escala errada"
    assert ws.cell(row=2, column=14).value == "12x36"


def test_acento_e_caixa_nao_reprovam_a_linha() -> None:
    planilha = build_workbook([linha(EVENTO_A)])
    preenchida = preencher(planilha, {2: ("Falso", "Tolerância Errada", "")})

    linhas, erros = parse(preenchida)

    assert erros == []
    assert linhas[0].cause == "wrong_tolerance"


# ---------------------------------------------------------------------------
# O que ela recusa — e o par legítimo de cada recusa
# ---------------------------------------------------------------------------


def test_falso_sem_causa_e_recusado_e_nada_e_gravado() -> None:
    planilha = build_workbook([linha(EVENTO_A), linha(EVENTO_B)])
    preenchida = preencher(planilha, {2: ("falso", "", ""), 3: ("falso", "erro do motor", "")})

    linhas, erros = parse(preenchida)

    assert [e.line for e in erros] == [2]
    assert "sem causa" in erros[0].message
    # A linha boa não é gravada por conta própria: o chamador não escreve nada
    # enquanto houver erro, e é por isso que ela volta aqui separada.
    assert [linha_.line for linha_ in linhas] == [3]


def test_verdadeiro_com_causa_e_recusado() -> None:
    planilha = build_workbook([linha(EVENTO_A)])
    preenchida = preencher(planilha, {2: ("verdadeiro", "erro do motor", "")})

    _, erros = parse(preenchida)

    assert [e.line for e in erros] == [2]
    assert "não leva causa" in erros[0].message


def test_palavra_fora_do_vocabulario_e_recusada() -> None:
    planilha = build_workbook([linha(EVENTO_A), linha(EVENTO_B)])
    preenchida = preencher(planilha, {2: ("talvez", "", ""), 3: ("falso", "escala torta", "")})

    linhas, erros = parse(preenchida)

    assert [e.line for e in erros] == [2, 3]
    assert linhas == []


def test_veredito_em_branco_nao_e_erro_e_sim_cobertura_faltando() -> None:
    planilha = build_workbook([linha(EVENTO_A), linha(EVENTO_B)])
    preenchida = preencher(planilha, {3: ("verdadeiro", "", "")})

    linhas, erros = parse(preenchida)

    assert erros == []
    assert [linha_.deviation_event_id for linha_ in linhas] == [EVENTO_B]


def test_id_invalido_e_recusado() -> None:
    planilha = build_workbook([linha(EVENTO_A)])
    wb = load_workbook(BytesIO(planilha))
    ws = wb["censo"]
    ws.cell(row=2, column=1, value="nao-e-uuid")
    ws.cell(row=2, column=12, value="verdadeiro")
    buffer = BytesIO()
    wb.save(buffer)

    linhas, erros = parse(buffer.getvalue())

    assert linhas == []
    assert "indicio_id inválido" in erros[0].message


def test_cabecalho_remontado_a_mao_recusa_o_arquivo_inteiro() -> None:
    """Coluna deslocada não dá erro de linha: dá veredito na pessoa errada."""
    planilha = build_workbook([linha(EVENTO_A)])
    wb = load_workbook(BytesIO(planilha))
    ws = wb["censo"]
    ws.cell(row=1, column=12, value="parecer")
    ws.cell(row=2, column=12, value="verdadeiro")
    buffer = BytesIO()
    wb.save(buffer)

    linhas, erros = parse(buffer.getvalue())

    assert linhas == []
    assert [e.line for e in erros] == [1]
    assert "cabeçalho" in erros[0].message


# ---------------------------------------------------------------------------
# A cobertura, que é o que separa uma medição de uma opinião
# ---------------------------------------------------------------------------


def medicao(events: int, judged: int, falsos: int) -> Measurement:
    total = Slice("__total__", events, judged, falsos)
    return Measurement(
        tenant_id=uuid4(),
        start=date(2026, 8, 29),
        end=date(2026, 9, 4),
        mode="shadow",
        total=total,
        by_type=[],
    )


def test_censo_incompleto_nao_passa_no_gate_nem_reprova() -> None:
    m = medicao(events=820, judged=40, falsos=0)

    assert m.total.rate == 0.0
    assert m.gate_passes is None


def test_censo_completo_no_limite_passa() -> None:
    assert medicao(events=100, judged=100, falsos=5).gate_passes is True


def test_censo_completo_acima_do_limite_reprova() -> None:
    assert medicao(events=100, judged=100, falsos=6).gate_passes is False


# ---------------------------------------------------------------------------
# A liberação recusa o que o schema também recusa — e antes de tocar o banco
# ---------------------------------------------------------------------------


def _release_com_medicao(monkeypatch: pytest.MonkeyPatch, m: Measurement):
    """`release` com a medição fixada e o banco PROIBIDO: abrir escopo é falha."""
    from operax.motor import adjudicacao

    async def medir_falso(*_: object, **__: object) -> Measurement:
        return m

    def escopo_proibido(*_: object, **__: object):
        raise AssertionError("release abriu conexão antes de a medição autorizar")

    monkeypatch.setattr(adjudicacao, "measure", medir_falso)
    monkeypatch.setattr(adjudicacao, "tenant_scope", escopo_proibido)
    return adjudicacao


@pytest.mark.anyio
async def test_liberacao_recusa_censo_incompleto_sem_tocar_o_banco(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    adjudicacao = _release_com_medicao(monkeypatch, medicao(events=820, judged=40, falsos=0))

    with pytest.raises(adjudicacao.ReleaseRefusedError, match="incompleto"):
        await adjudicacao.release(
            object(), date(2026, 8, 12), date(2026, 9, 10), mode="shadow", author="x", note=None
        )


@pytest.mark.anyio
async def test_liberacao_recusa_taxa_acima_do_teto_sem_tocar_o_banco(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    adjudicacao = _release_com_medicao(monkeypatch, medicao(events=100, judged=100, falsos=6))

    with pytest.raises(adjudicacao.ReleaseRefusedError, match="acima do teto"):
        await adjudicacao.release(
            object(), date(2026, 8, 12), date(2026, 9, 10), mode="shadow", author="x", note=None
        )
