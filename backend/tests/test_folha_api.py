"""A folha inteira: modelo → preenchimento → preview → confirmação → reenvio.

Nada aqui abre banco, e nada aqui finge SQL: `tenant_scope` e `user_scope` são as
duas costuras, e o que entra no lugar delas é um cursor de mentira que executa as
instruções de verdade — inclusive `bind_tenant`, que recusa qualquer statement
sem filtro de tenant.

O caminho é um teste só de propósito. O valor da entrega não está em nenhuma
etapa isolada: está no arquivo voltar preenchido com o que entrou, o reenvio
**substituir** a competência em vez de somar, e a folha com uma linha errada não
ser importada de jeito nenhum — que é a diferença de comportamento em relação ao
import de RH, e a única que muda um número que alguém vai comparar com holerite.
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any
from uuid import UUID, uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from operax.core.tenant import bind_tenant
from operax.dp import rubricas as dp_rubricas
from operax.imports import payroll
from operax.imports import repository as folha_repository
from operax.rh import repository as rh_repository
from operax.rh.workbook import META_SHEET
from server.main import app
from server.routers import folha

EMPRESA = UUID("55555555-5555-4555-8555-555555555555")
UNIDADE = UUID("66666666-6666-4666-8666-666666666666")

PESSOAS = [
    {
        "employee_id": uuid4(),
        "registration_number": f"100{n}",
        "name": nome,
        "company_id": EMPRESA,
        "unit_id": UNIDADE,
    }
    for n, nome in enumerate(["Ana Personagem", "Bruno Personagem"], start=1)
]

#: Só um dos códigos usados está curado. O outro é o caso que interessa: entra,
#: conta no total, e volta na lista da curadoria.
#:
#: ⛔ "CURADO" É CLASSIFICADO **E** CONFERIDO, e é por isso que isto é uma linha
#: com colunas em vez de uma lista de códigos. `dp_payroll_code_map` semeia uma
#: linha por código da FOLHA com `category` nula: depois dela, "existe linha no
#: mapa" deixou de significar "tem categoria", e um dublê que guardasse só o
#: código não teria como distinguir os dois estados — que é exatamente o defeito.
CURADORIA = [{"code": "0050", "category": "overtime", "validated": True}]


class FakeCursor:
    """Um banco de mentira endereçado por statement, não por ordem de chamada."""

    def __init__(self, estado: FakeDB, context: Any, checar_tenant: bool) -> None:
        self._estado = estado
        self._context = context
        self._checar_tenant = checar_tenant
        self._resultado: Any = None

    async def execute(self, statement: str, params: Any = None) -> None:
        if self._checar_tenant:
            params = bind_tenant(statement, params, self._context)
        params = params or {}
        self._estado.statements.append((" ".join(statement.split()), dict(params)))
        self._resultado = self._estado.responder(statement, params)

    async def fetchone(self) -> Any:
        return self._resultado[0] if self._resultado else None

    async def fetchall(self) -> list[dict[str, Any]]:
        return list(self._resultado or [])


class FakeScope:
    def __init__(self, cursor: FakeCursor) -> None:
        self._cursor = cursor

    async def __aenter__(self) -> FakeCursor:
        return self._cursor

    async def __aexit__(self, *exc: object) -> None:
        return None


class FakeDB:
    """Estado do cliente: pessoas, competências, lançamentos e a trilha."""

    def __init__(self, *, admin: bool = True, dominios: bool = True) -> None:
        self.pessoas = [dict(p) for p in PESSOAS]
        self.curadoria: list[dict[str, Any]] = []
        for linha in CURADORIA:
            self.semear_curadoria(**linha)
        self.periods: dict[tuple[int, int], dict[str, Any]] = {}
        self.entries: list[dict[str, Any]] = []
        self.imports: dict[UUID, dict[str, Any]] = {}
        self.audit: list[dict[str, Any]] = []
        self.statements: list[tuple[str, dict[str, Any]]] = []
        self.admin = admin
        self.dominios = dominios

    def semear_curadoria(
        self,
        code: str,
        *,
        category: str | None,
        validated: bool,
        label: str | None = None,
        nature: str | None = None,
    ) -> None:
        """Uma linha de `app.payroll_event_map`, no estado em que ela de fato existe.

        `category=None, validated=False` é o estado em que a SEMENTE entrega todo
        código da folha — e o banco proíbe o inverso
        (`payroll_event_map_validated_has_category`), então validar sem
        classificar não é semeável aqui de propósito.
        """
        assert not validated or category is not None, (
            "o banco recusa validated_at sem category (payroll_event_map_validated_has_category)"
        )
        self.curadoria.append(
            {
                "code": code,
                "label": label,
                "nature": nature,
                "category": category,
                "validated_at": datetime(2026, 9, 1, 12, 0) if validated else None,
            }
        )

    def abrir_competencia(self, year: int, month: int, status: str) -> UUID:
        period_id = uuid4()
        self.periods[(year, month)] = {"id": period_id, "status": status}
        return period_id

    def responder(self, statement: str, params: dict[str, Any]) -> Any:
        sql = " ".join(statement.split())

        if "util.is_admin" in sql:
            return [
                {
                    "admin": self.admin,
                    "pii": self.dominios,
                    "compensation": self.dominios,
                    "health": self.dominios,
                }
            ]
        if "from app.payroll_entry pe" in sql:
            return self._entries_da_competencia(params["year"], params["month"])
        if "from app.employee e" in sql:
            return [dict(p) for p in self.pessoas]
        if "from app.payroll_event_map" in sql:
            return self._lista_de_rubricas(params)
        if "from app.payroll_period p" in sql:
            return self._estado_da_competencia(params["year"], params["month"])
        if "insert into app.payroll_period" in sql:
            return self._garantir_competencia(params["year"], params["month"])
        if "delete from app.payroll_entry" in sql:
            apagadas = [e for e in self.entries if e["period_id"] == params["period_id"]]
            self.entries = [e for e in self.entries if e["period_id"] != params["period_id"]]
            return [{"id": uuid4()} for _ in apagadas]
        if "insert into app.payroll_entry" in sql:
            self.entries.append(dict(params))
            return []
        if "insert into app.file_import" in sql:
            self.imports[params["import_id"]] = {
                "id": params["import_id"],
                "type": params["type"],
                "storage_path": params["storage_path"],
                "file_name": params["file_name"],
                "layout_version": params["layout_version"],
                "status": "received",
                "payroll_period_id": None,
                "rows_total": None,
                "rows_ok": None,
                "rows_error": None,
                "report": None,
            }
            return []
        if "update app.file_import set status" in sql:
            self.imports[params["import_id"]] |= {
                "status": params["status"],
                "rows_total": params["rows_total"],
                "rows_ok": params["rows_ok"],
                "rows_error": params["rows_error"],
                "report": params["report"],
            }
            return []
        if "update app.file_import set payroll_period_id" in sql:
            self.imports[params["import_id"]]["payroll_period_id"] = params["period_id"]
            return []
        if "from app.file_import" in sql:
            registro = self.imports.get(params["import_id"])
            return [registro] if registro else []
        if "insert into app.audit_log" in sql:
            self.audit.append(dict(params))
            return []
        return []

    # -- curadoria de rubrica -------------------------------------------------
    def _lista_de_rubricas(self, params: dict[str, Any]) -> list[dict[str, Any]]:
        """As duas fontes de `rubricas._LIST_SQL`: a folha importada e o mapa curado.

        ⛔ O DUBLÊ NÃO SEPARA O CURADO DO NÃO CURADO. Ele devolve as duas metades
        como a consulta as devolve, com `category` e `validated_at` do jeito que
        estão; quem separa é `rubricas.read_curation`, que é o código sob teste.
        Um fake que filtrasse por `category is not null` aqui daria verde à
        pergunta errada e esconderia o defeito.
        """
        curados = {linha["code"]: linha for linha in self.curadoria}
        da_folha = {linha["code"]: linha for linha in self.entries}
        linhas = []
        for code in sorted(set(curados) | set(da_folha)):
            curado = curados.get(code, {})
            entrada = da_folha.get(code, {})
            linhas.append(
                {
                    "code": code,
                    "label": curado.get("label") or entrada.get("description"),
                    "nature": curado.get("nature") or entrada.get("nature"),
                    "category": curado.get("category"),
                    "validated_at": curado.get("validated_at"),
                    "in_payroll": code in da_folha,
                }
            )
        if params.get("code") is not None:
            linhas = [linha for linha in linhas if linha["code"] == params["code"]]
        return linhas

    # -- competência ---------------------------------------------------------
    def _estado_da_competencia(self, year: int, month: int) -> list[dict[str, Any]]:
        periodo = self.periods.get((year, month))
        if periodo is None:
            return []
        return [
            {
                "id": periodo["id"],
                "status": periodo["status"],
                "entries": len([e for e in self.entries if e["period_id"] == periodo["id"]]),
                "imported_at": periodo.get("imported_at"),
            }
        ]

    def _garantir_competencia(self, year: int, month: int) -> list[dict[str, Any]]:
        periodo = self.periods.get((year, month))
        if periodo is None:
            return [{"id": self.abrir_competencia(year, month, "importada")}]
        if periodo["status"] == "fechada":
            # O `where` do `on conflict do update` não casa: nenhuma linha volta.
            return []
        periodo["status"] = "importada"
        return [{"id": periodo["id"]}]

    def _entries_da_competencia(self, year: int, month: int) -> list[dict[str, Any]]:
        periodo = self.periods.get((year, month))
        if periodo is None:
            return []
        por_id = {p["employee_id"]: p for p in self.pessoas}
        linhas = []
        for entrada in self.entries:
            if entrada["period_id"] != periodo["id"]:
                continue
            pessoa = por_id[entrada["employee_id"]]
            linhas.append(
                {
                    "employee_code": pessoa["registration_number"],
                    "employee_name": pessoa["name"],
                    "code": entrada["code"],
                    "description": entrada["description"],
                    "nature": entrada["nature"],
                    "reference": entrada["reference"],
                    "amount": entrada["amount"],
                }
            )
        return sorted(linhas, key=lambda linha: (linha["employee_code"], linha["code"]))


class FakeStore:
    """Storage em memória. O confirm relê daqui, como releria do Supabase."""

    def __init__(self) -> None:
        self.objetos: dict[str, bytes] = {}

    async def put(self, path: str, data: bytes, content_type: str) -> None:
        self.objetos[path] = data

    async def get(self, path: str) -> bytes:
        return self.objetos[path]


@pytest.fixture
def db(monkeypatch: pytest.MonkeyPatch):
    def instalar(**kwargs: Any) -> FakeDB:
        estado = FakeDB(**kwargs)
        # `dp_rubricas` entra porque `fetch_mapped_codes` pergunta a ele quem está
        # curado — a porta única de categoria é dele, não do import. Ele só lê sob
        # `tenant_scope`, e por isso não aparece na segunda lista.
        for modulo in (folha_repository, rh_repository, dp_rubricas):
            monkeypatch.setattr(
                modulo,
                "tenant_scope",
                lambda context, schema="app": FakeScope(FakeCursor(estado, context, True)),
            )
        for modulo in (folha_repository, rh_repository):
            monkeypatch.setattr(
                modulo,
                "user_scope",
                lambda context, schema="app": FakeScope(FakeCursor(estado, context, False)),
            )
        return estado

    return instalar


@pytest.fixture
def store() -> FakeStore:
    memoria = FakeStore()
    app.dependency_overrides[folha.get_store] = lambda: memoria
    return memoria


# ---------------------------------------------------------------------------
# Utilidades da planilha
# ---------------------------------------------------------------------------
def auth(issue_token) -> dict[str, str]:
    return {"Authorization": f"Bearer {issue_token()}"}


def baixar(client: TestClient, issue_token, *, ano: int = 2026, mes: int = 8):
    return client.get(f"/folha/template?ano={ano}&mes={mes}", headers=auth(issue_token))


def preencher(conteudo: bytes, linhas: list[dict[str, Any]]) -> bytes:
    """Escreve as linhas pelo rótulo da coluna, como quem preenche o modelo."""
    wb = load_workbook(BytesIO(conteudo))
    ws = wb[payroll.SHEET_TITLE]
    rotulos = {celula.value: celula.column for celula in next(ws.iter_rows(min_row=1, max_row=1))}
    for deslocamento, valores in enumerate(linhas):
        for rotulo, valor in valores.items():
            ws.cell(row=2 + deslocamento, column=rotulos[rotulo]).value = valor
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def enviar(client: TestClient, issue_token, conteudo: bytes):
    return client.post(
        "/folha/imports",
        headers=auth(issue_token),
        files={"arquivo": ("folha.xlsx", conteudo, payroll.CONTENT_TYPE)},
    )


def linha(matricula: str, codigo: str, valor: str, *, natureza: str = "Provento") -> dict[str, Any]:
    return {
        "Matrícula": matricula,
        "Colaborador": "quem preencheu escreveu",
        "Código do evento": codigo,
        "Descrição do evento": "Salário base",
        "Natureza": natureza,
        "Valor": valor,
    }


# ---------------------------------------------------------------------------
# O caminho inteiro
# ---------------------------------------------------------------------------
def test_do_modelo_ate_a_substituicao_da_competencia(client: TestClient, issue_token, db, store):
    estado = db()

    modelo = baixar(client, issue_token)
    assert modelo.status_code == 200
    assert modelo.headers["content-type"] == payroll.CONTENT_TYPE
    assert "folha-2026-08.xlsx" in modelo.headers["content-disposition"]
    # O download em si é registrado: o arquivo leva o dado embora.
    assert [a["action"] for a in estado.audit] == ["export"]

    preenchido = preencher(
        modelo.content,
        [
            linha("1001", "0050", "2.500,00"),
            # Formato brasileiro e código fora do mapa: os dois casos que o
            # arquivo real traz e que não podem virar recusa.
            linha("1001", "H_EXTRA_60", "312,45"),
            linha("1002", "0050", "1.900,00"),
        ],
    )

    preview = enviar(client, issue_token, preenchido)
    assert preview.status_code == 201
    corpo = preview.json()
    assert corpo["period"] == "2026-08"
    assert corpo["status"] == "validating"
    assert corpo["counts"] == {"total": 3, "ok": 3, "error": 0}
    assert corpo["unmapped_codes"] == ["H_EXTRA_60"]
    assert corpo["replaces"] is None
    # Só a linha que tem o que dizer aparece — e o que ela tem é aviso, não erro.
    assert [(linha_["line"], linha_["warnings"][0]["code"]) for linha_ in corpo["lines"]] == [
        (3, "codigo_sem_categoria")
    ]

    # Preview não grava: nem lançamento, nem competência, nem trilha nova.
    assert estado.entries == []
    assert estado.periods == {}
    assert [a["action"] for a in estado.audit] == ["export"]

    import_id = corpo["import_id"]
    confirmado = client.post(f"/folha/imports/{import_id}/confirm", headers=auth(issue_token))
    assert confirmado.status_code == 200
    resultado = confirmado.json()
    assert (resultado["applied"], resultado["replaced"]) == (3, 0)
    assert resultado["status"] == "processed"

    assert len(estado.entries) == 3
    gravada = estado.entries[0]
    assert gravada["amount"] == Decimal("2500.00")
    assert gravada["nature"] == "earning"
    assert gravada["source"] == "spreadsheet"
    assert gravada["import_id"] == UUID(import_id)
    # Empresa e unidade vêm do colaborador, nunca do departamento — regra 5.
    assert (gravada["company_id"], gravada["unit_id"]) == (EMPRESA, UNIDADE)

    # A trilha é do arquivo, não da linha: um `insert` para as três.
    assert [a["action"] for a in estado.audit] == ["export", "insert"]
    trilha = estado.audit[-1]["depois"].obj
    assert trilha["period"] == "2026-08"
    assert trilha["rows"] == 3
    assert trilha["amount_total"] == Decimal("4712.45")
    assert trilha["_origem"]["file_import_id"] == import_id

    guardado = estado.imports[UUID(import_id)]
    assert (guardado["status"], guardado["rows_ok"], guardado["rows_error"]) == ("processed", 3, 0)
    assert guardado["payroll_period_id"] == estado.periods[(2026, 8)]["id"]
    assert estado.periods[(2026, 8)]["status"] == "importada"

    # --- o reenvio da competência corrigida --------------------------------
    segundo_modelo = baixar(client, issue_token)
    ws = load_workbook(BytesIO(segundo_modelo.content))[payroll.SHEET_TITLE]
    # O modelo volta preenchido com o que entrou, e a natureza volta em
    # português: corrigir é mexer numa célula, não redigitar a folha.
    assert [c.value for c in ws["A"][1:]] == ["1001", "1001", "1002"]
    assert [c.value for c in ws["E"][1:]] == ["Provento", "Provento", "Provento"]

    corrigido = preencher(
        segundo_modelo.content,
        [
            linha("1001", "0050", "2.500,00"),
            linha("1001", "H_EXTRA_60", "512,45"),
            linha("1002", "0050", "1.900,00"),
        ],
    )
    segundo_preview = enviar(client, issue_token, corrigido)
    assert segundo_preview.json()["replaces"]["entries"] == 3

    segundo_id = segundo_preview.json()["import_id"]
    fechamento = client.post(f"/folha/imports/{segundo_id}/confirm", headers=auth(issue_token))
    assert fechamento.status_code == 200
    assert (fechamento.json()["applied"], fechamento.json()["replaced"]) == (3, 3)

    # Substituiu, não somou: continuam três linhas, e o valor corrigido é o que
    # está lá.
    assert len(estado.entries) == 3
    assert sorted(e["amount"] for e in estado.entries) == [
        Decimal("512.45"),
        Decimal("1900.00"),
        Decimal("2500.00"),
    ]
    assert [a["action"] for a in estado.audit[-2:]] == ["delete", "insert"]
    assert estado.audit[-2]["antes"].obj == {"period": "2026-08", "rows": 3}


def test_codigo_semeado_no_mapa_e_sem_categoria_continua_pendente(
    client: TestClient, issue_token, db, store
):
    """⛔ "EXISTE LINHA NO MAPA" NÃO É "TEM CATEGORIA" — e o import perguntava a errada.

    `20260907182521_dp_payroll_code_map.sql` semeia UMA LINHA POR CÓDIGO DA FOLHA
    com `category` nula: é assim que a lista chega pronta para a contabilidade
    conferir. Enquanto `fetch_mapped_codes` perguntava
    `select code from app.payroll_event_map`, todo código semeado passava por
    curado — `unmapped_codes` voltava VAZIO com a curadoria inteira por fazer, e a
    tela de import dizia "nenhuma pendência" enquanto a de rubricas listava N.

    AS DUAS CONDIÇÕES QUE SOZINHAS JÁ DARIAM VERDE, E POR QUE NÃO SÃO ELAS:
      · o dublê não devolver a linha semeada — excluída pela asserção sobre
        `estado.curadoria`, que é o estado inteiro que o responder devolve;
      · o import chamar TODO código de não curado — excluída pela lista exata,
        que traz `H_EXTRA_60` e **não** traz `0050`, curado no mesmo cenário.
    """
    estado = db()
    # O estado exato em que a semente entrega: a linha existe e não tem categoria.
    estado.semear_curadoria("H_EXTRA_60", category=None, validated=False)
    assert [(linha_["code"], linha_["category"]) for linha_ in estado.curadoria] == [
        ("0050", "overtime"),
        ("H_EXTRA_60", None),
    ]

    preenchido = preencher(
        baixar(client, issue_token).content,
        [linha("1001", "0050", "2.500,00"), linha("1001", "H_EXTRA_60", "312,45")],
    )
    corpo = enviar(client, issue_token, preenchido).json()

    assert corpo["counts"] == {"total": 2, "ok": 2, "error": 0}
    assert corpo["unmapped_codes"] == ["H_EXTRA_60"]
    assert [(linha_["line"], linha_["warnings"][0]["code"]) for linha_ in corpo["lines"]] == [
        (3, "codigo_sem_categoria")
    ]


# ---------------------------------------------------------------------------
# Recusas
# ---------------------------------------------------------------------------
def test_linha_em_erro_nao_e_confirmavel(client: TestClient, issue_token, db, store):
    estado = db()
    preenchido = preencher(
        baixar(client, issue_token).content,
        [
            linha("1001", "0050", "2.500,00"),
            # Matrícula que não existe: a folha não entra pela metade, então esta
            # linha derruba o arquivo inteiro — e não só a si mesma.
            linha("9999", "0050", "1.000,00"),
        ],
    )

    preview = enviar(client, issue_token, preenchido)
    corpo = preview.json()
    assert corpo["status"] == "validation_error"
    assert corpo["counts"] == {"total": 2, "ok": 1, "error": 1}
    assert corpo["lines"][0]["errors"][0]["code"] == "colaborador_desconhecido"

    confirmado = client.post(
        f"/folha/imports/{corpo['import_id']}/confirm", headers=auth(issue_token)
    )
    assert confirmado.status_code == 409
    assert "validation_error" in confirmado.json()["detail"]
    assert estado.entries == []
    assert [a["action"] for a in estado.audit] == ["export"]


def test_competencia_fechada_recusa_o_arquivo_e_guarda_o_motivo(
    client: TestClient, issue_token, db, store
):
    estado = db()
    estado.abrir_competencia(2026, 8, "fechada")

    modelo = baixar(client, issue_token)
    resposta = enviar(
        client, issue_token, preencher(modelo.content, [linha("1001", "0050", "2.500,00")])
    )
    assert resposta.status_code == 409
    assert "fechada" in resposta.json()["detail"]

    # O arquivo recusado continua guardado, e o registro dele diz por quê.
    assert len(store.objetos) == 1
    guardado = next(iter(estado.imports.values()))
    assert guardado["status"] == "validation_error"
    assert guardado["report"].obj["file_error"]["code"] == "competencia_fechada"
    assert estado.entries == []


def test_papel_sem_permissao_nao_envia_folha(client: TestClient, issue_token, db, store):
    db(admin=False)
    resposta = client.post(
        "/folha/imports",
        headers=auth(issue_token),
        files={"arquivo": ("folha.xlsx", b"nao chega a ser lido", payroll.CONTENT_TYPE)},
    )
    assert resposta.status_code == 403
    assert store.objetos == {}


def test_sem_dominio_de_remuneracao_nao_baixa_o_modelo(client: TestClient, issue_token, db, store):
    db(dominios=False)
    resposta = baixar(client, issue_token)
    assert resposta.status_code == 403
    assert "remuneração" in resposta.json()["detail"]


# ---------------------------------------------------------------------------
# A guarda que não se contorna
# ---------------------------------------------------------------------------
async def test_gravacao_recusa_arquivo_com_erro_mesmo_sem_passar_pelo_endpoint(db):
    """A recusa da folha parcial vive na gravação, não só na rota.

    O endpoint já barra pelo estado do import. Esta é a segunda tranca: qualquer
    caminho novo até `apply_payroll` — um reprocessamento, uma tarefa agendada —
    encontra a mesma recusa, porque a soma errada é a mesma.
    """
    from operax.core.tenant import TenantContext, UserRole
    from operax.rh.validators import LineError

    estado = db()
    contexto = TenantContext(
        tenant_id=UUID("22222222-2222-4222-8222-222222222222"),
        user_id=UUID("11111111-1111-4111-8111-111111111111"),
        role=UserRole.HR,
    )
    reprovada = payroll.LineOutcome(
        line=2,
        employee_id=None,
        values={},
        errors=(LineError("colaborador_desconhecido", "não existe", "employee_code"),),
    )

    with pytest.raises(folha_repository.PartialPayrollError):
        await folha_repository.apply_payroll(
            contexto,
            import_id=uuid4(),
            year=2026,
            month=8,
            outcomes=(reprovada,),
            employees={},
            rows_total=1,
            rows_ok=0,
            rows_error=1,
            report={},
        )
    assert estado.entries == []
    assert estado.periods == {}


def test_modelo_da_competencia_carrega_a_aba_de_controle(
    client: TestClient, issue_token, db, store
):
    """O modelo baixado é o que o parser reconhece — as duas pontas do mesmo arquivo."""
    db()
    conteudo = baixar(client, issue_token, ano=2026, mes=7).content
    wb = load_workbook(BytesIO(conteudo))
    assert META_SHEET in wb.sheetnames

    lido = payroll.parse_upload(conteudo, tenant_id=UUID("22222222-2222-4222-8222-222222222222"))
    assert (lido.year, lido.month) == (2026, 7)
    assert lido.rows == ()
