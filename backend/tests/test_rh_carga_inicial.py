"""O ensaio da carga: a planilha bagunçada entra, e nada sai pela porta dos fundos.

O critério de sucesso do PRD — "toda linha da planilha original tem destino no
sistema **ou** entrada no relatório de descarte com motivo" — é uma frase até
alguém somar as duas colunas. Aqui ele é uma asserção.

O resto da suíte cerca as três formas de o conversor errar em silêncio: escrever
no arquivo errado (a identidade do modelo), escrever a pessoa errada (a chave), e
escrever o que a regra proíbe (CID, restrição, conta bancária).
"""

from __future__ import annotations

import inspect
from datetime import date
from pathlib import Path
from uuid import UUID

import pytest

from operax.rh import carga_inicial
from operax.rh.carga_inicial import (
    SHEETS,
    ConverterError,
    Purpose,
    Resultado,
    converter,
    normalize,
)
from operax.rh.ownership import ENUMS, tables
from operax.rh.templates import TEMPLATES, get_template
from operax.rh.workbook import parse
from tests.fixtures.planilha_cliente import (
    FORA_DO_QUADRO,
    PESSOAS,
    TENANT,
    escrever_modelos,
    escrever_planilha,
)


@pytest.fixture(scope="module")
def ensaio(tmp_path_factory: pytest.TempPathFactory) -> tuple[Resultado, Path]:
    """A corrida inteira, uma vez: ler, converter, emitir."""
    raiz = tmp_path_factory.mktemp("carga")
    planilha = escrever_planilha(raiz / "planilha-do-cliente.xlsx")
    modelos = escrever_modelos(raiz / "modelos")
    saida = raiz / "saida"
    return converter(planilha=planilha, modelos=modelos, saida=saida), saida


def _texto_de_tudo(saida: Path, *, com_relatorio: bool = True) -> str:
    """Tudo o que o conversor gravou, como texto — inclusive dentro dos .xlsx."""
    from openpyxl import load_workbook

    pedacos: list[str] = []
    for caminho in sorted(saida.rglob("*")):
        if not com_relatorio and caminho.suffix == ".md":
            continue
        if caminho.suffix == ".xlsx":
            wb = load_workbook(caminho)
            for ws in wb.worksheets:
                for linha in ws.iter_rows():
                    pedacos += [str(c.value) for c in linha if c.value is not None]
        elif caminho.is_file():
            pedacos.append(caminho.read_text(encoding="utf-8"))
    return "\n".join(pedacos)


# ---------------------------------------------------------------------------
# O gate
# ---------------------------------------------------------------------------
def test_o_ensaio_fecha_em_cem_por_cento(ensaio: tuple[Resultado, Path]) -> None:
    resultado, _ = ensaio

    assert resultado.lidas > 0
    assert resultado.fecha(), (
        f"{resultado.lidas} lidas contra "
        f"{resultado.destino + resultado.parqueadas + resultado.conferidas + resultado.descartadas}"
        " distribuídas"
    )
    # E cada descarte tem motivo escrito: "descartada" sem motivo é linha perdida
    # com contabilidade em dia.
    assert all(descarte.reason for descarte in resultado.descartes)


def test_todo_descarte_tem_motivo_no_relatorio(ensaio: tuple[Resultado, Path]) -> None:
    resultado, saida = ensaio
    relatorio = (saida / "RELATORIO-DE-DESCARTE.md").read_text(encoding="utf-8")

    for descarte in resultado.descartes:
        assert descarte.reason in relatorio


# ---------------------------------------------------------------------------
# Identidade do arquivo
# ---------------------------------------------------------------------------
def test_o_modelo_preenchido_volta_pelo_parser_sem_recusa(ensaio: tuple[Resultado, Path]) -> None:
    """A prova de que o conversor preenche em vez de recriar.

    `parse` recusa o arquivo inteiro se o hash do cabeçalho, o tenant ou a versão
    de layout não baterem. Passar por ele é a única forma de afirmar que o
    arquivo emitido aqui é aceito lá.
    """
    _, saida = ensaio

    for tipo in TEMPLATES:
        caminho = next(saida.glob(f"*-{tipo}.xlsx"))
        template = get_template(tipo)
        assert template is not None
        lido = parse(caminho.read_bytes(), template=template, tenant_id=TENANT)

        assert lido.layout_version == template.layout_version
        assert len(lido.rows) == len(PESSOAS)


def test_o_valor_convertido_chega_na_linha_da_pessoa_certa(ensaio: tuple[Resultado, Path]) -> None:
    _, saida = ensaio
    template = get_template("hr_link")
    assert template is not None
    lido = parse(
        next(saida.glob("*-hr_link.xlsx")).read_bytes(), template=template, tenant_id=TENANT
    )

    por_matricula = {str(r.values["registration_number"]): r.values for r in lido.rows}
    # A ordem da planilha do cliente é a mesma do modelo aqui, mas o casamento é
    # por matrícula: se fosse por posição, uma linha a mais na origem deslocaria
    # o ID RH de todo mundo a partir dela.
    assert por_matricula[PESSOAS[0][0]]["hr_code"] == "RH-001"
    assert por_matricula[PESSOAS[5][0]]["hr_code"] == "RH-006"


def test_a_matricula_que_o_sistema_nao_conhece_nao_entra(ensaio: tuple[Resultado, Path]) -> None:
    resultado, saida = ensaio

    assert FORA_DO_QUADRO not in _texto_de_tudo(saida).replace(
        "matrícula fora do quadro", ""
    ) or FORA_DO_QUADRO in "".join(e for d in resultado.descartes for e in d.examples)
    motivos = [d.reason for d in resultado.descartes if FORA_DO_QUADRO in "".join(d.examples)]
    assert any("quadro ativo" in m for m in motivos)


def test_linha_sem_matricula_e_descarte_com_motivo(ensaio: tuple[Resultado, Path]) -> None:
    resultado, _ = ensaio

    assert any("sem matrícula" in d.reason for d in resultado.descartes)


# ---------------------------------------------------------------------------
# A bagunça do arquivo
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    ("aba", "linha"),
    [("DADOS FUNCIONÁRIOS", 6), ("DESLIGADOS", 2), ("FÉRIAS", 8), ("FALTAS_ATESTADOS", 1)],
)
def test_o_cabecalho_e_encontrado_onde_ele_esta(
    ensaio: tuple[Resultado, Path], aba: str, linha: int
) -> None:
    """Título, legenda e linha em branco em cima do dado não deslocam a leitura."""
    resultado, _ = ensaio
    leitura = next(leitura for leitura in resultado.leituras if leitura.title == aba)

    assert leitura.header_row == linha


def test_a_aba_oculta_e_lida_como_qualquer_outra(ensaio: tuple[Resultado, Path]) -> None:
    resultado, saida = ensaio
    leitura = next(leitura for leitura in resultado.leituras if leitura.title == "CNH")

    assert leitura.hidden is True
    assert len(leitura.rows) == 2
    assert (saida / "parqueado" / "document.csv").exists()


def test_formula_quebrada_nunca_vira_conteudo(ensaio: tuple[Resultado, Path]) -> None:
    resultado, saida = ensaio

    assert sum(leitura.broken_cells for leitura in resultado.leituras) > 0
    # Fora o relatório, que cita o texto do erro justamente para dizer que o
    # ignorou.
    assert "#REF!" not in _texto_de_tudo(saida, com_relatorio=False)


def test_aba_duplicada_e_descartada_inteira(ensaio: tuple[Resultado, Path]) -> None:
    """A cópia não reimporta ninguém: importar as duas dobraria cada pessoa."""
    resultado, saida = ensaio
    duplicada = next(leitura for leitura in resultado.leituras if leitura.title.endswith("(2)"))

    assert duplicada.duplicate_of is not None
    assert duplicada.rows == []
    assert any("duplicada" in d.reason for d in resultado.descartes)

    template = get_template("hr_link")
    assert template is not None
    lido = parse(
        next(saida.glob("*-hr_link.xlsx")).read_bytes(), template=template, tenant_id=TENANT
    )
    matriculas = [str(r.values["registration_number"]) for r in lido.rows]
    assert len(matriculas) == len(set(matriculas))


def test_aba_fora_do_mapa_e_contada_e_nomeada(ensaio: tuple[Resultado, Path]) -> None:
    resultado, _ = ensaio

    fora = next(leitura for leitura in resultado.leituras if leitura.title == "ANIVERSARIANTES")
    assert fora.sheet is None
    assert fora.unread == 2  # duas pessoas; a primeira linha é dada como cabeçalho
    assert any("fora do mapa" in d.reason for d in resultado.descartes)


# ---------------------------------------------------------------------------
# O que a regra proíbe
# ---------------------------------------------------------------------------
def test_nenhum_cid_e_nenhuma_restricao_saem_do_conversor(ensaio: tuple[Resultado, Path]) -> None:
    """Regra 10, provada no arquivo e não no desenho.

    Os códigos estão na fixture de propósito: o teste só vale se o dado proibido
    tiver estado ali para ser recusado.
    """
    _, saida = ensaio
    tudo = _texto_de_tudo(saida)

    for proibido in ("F32", "M54", "J11", "K52", "não pode subir escada"):
        assert proibido not in tudo


def test_o_relatorio_nomeia_a_coluna_recusada_em_vez_de_omiti_la(
    ensaio: tuple[Resultado, Path],
) -> None:
    resultado, saida = ensaio
    relatorio = (saida / "RELATORIO-DE-DESCARTE.md").read_text(encoding="utf-8")
    recusadas = {coluna for leitura in resultado.leituras for coluna, _ in leitura.refused_found}

    # A prova de LGPD é dizer "estava aqui e não entrou", não ficar em silêncio.
    assert {"CID", "RESTRICAO", "BANCO", "CONTA", "PIX"} <= recusadas
    assert "regra 10" in relatorio


def test_dado_bancario_nao_sai_do_conversor(ensaio: tuple[Resultado, Path]) -> None:
    _, saida = ensaio
    tudo = _texto_de_tudo(saida)

    assert "56789-0" not in tudo
    assert "@exemplo.invalido" not in tudo


# ---------------------------------------------------------------------------
# As conversões declaradas no sprint
# ---------------------------------------------------------------------------
def test_desligados_vira_conferencia_e_nao_escrita(ensaio: tuple[Resultado, Path]) -> None:
    """`employee.status` é do sync: a aba serve para apontar a divergência.

    Duas das quatro linhas da fixture nomeiam gente que o ponto ainda traz ativa
    — essas são divergência. As outras duas já saíram do quadro, e conferir uma
    aba que fecha não pode produzir alarme.
    """
    resultado, saida = ensaio

    assert resultado.conferidas == 4
    assert len(resultado.divergencias) == 2
    assert all("quadro ativo" in d for d in resultado.divergencias)
    # E nada da aba virou dado: a data de demissão da planilha aparece no
    # relatório, para o implantador levar ao Secullum, e em lugar nenhum que o
    # sistema vá ler. `status` continua sendo do sync.
    assert "30/06/2026" in (saida / "RELATORIO-DE-DESCARTE.md").read_text(encoding="utf-8")
    assert "2026-06-30" not in _texto_de_tudo(saida, com_relatorio=False)


def test_ferias_em_matriz_vira_formato_longo(ensaio: tuple[Resultado, Path]) -> None:
    resultado, saida = ensaio
    linhas = (saida / "parqueado" / "leave_period.csv").read_text(encoding="utf-8").splitlines()
    ferias = [linha for linha in linhas if "vacation" in linha]

    # Três pessoas, quatro períodos completos: um ano por par de colunas.
    assert len(ferias) == 4
    assert f"{PESSOAS[0][0]},vacation,2024-07-01,2024-07-30" in ferias
    # O ano com só um lado não vira período pela metade — vira pergunta.
    assert any("período de 2026" in n.reason for n in resultado.notas)


def test_aso_em_portugues_vira_o_vocabulario_do_banco(ensaio: tuple[Resultado, Path]) -> None:
    _, saida = ensaio
    template = get_template("hr_exam")
    assert template is not None
    lido = parse(
        next(saida.glob("*-hr_exam.xlsx")).read_bytes(), template=template, tenant_id=TENANT
    )
    por_matricula = {str(r.values["registration_number"]): r.values for r in lido.rows}

    assert por_matricula[PESSOAS[0][0]]["type"] == "periodic"
    assert por_matricula[PESSOAS[1][0]]["type"] == "pre_employment"
    # "APTO COM RESTRIÇÃO" é aptidão e cabe; a descrição da restrição não.
    assert por_matricula[PESSOAS[2][0]]["result"] == "fit_with_restriction"
    assert "escada" not in _texto_de_tudo(saida)


def test_o_historico_de_aso_e_parqueado_em_vez_de_sobrescrever(
    ensaio: tuple[Resultado, Path],
) -> None:
    """A pessoa tem dois exames na aba, e o modelo carrega um por pessoa.

    Sem a redução, as duas linhas escreveriam na mesma célula e a última lida
    ganharia — em silêncio, pela ordem da planilha. O que vence é a data maior, e
    o anterior sai convertido no .csv em vez de sumir.
    """
    resultado, saida = ensaio
    assert resultado.historico.get("VENCIMENTO ASO") == 1

    historico = (
        (saida / "parqueado" / "occupational_exam.csv").read_text(encoding="utf-8").splitlines()
    )
    anterior = f"{PESSOAS[0][0]},pre_employment,2024-08-15"
    assert any(linha.startswith(anterior) for linha in historico)

    template = get_template("hr_exam")
    assert template is not None
    lido = parse(
        next(saida.glob("*-hr_exam.xlsx")).read_bytes(), template=template, tenant_id=TENANT
    )
    por_matricula = {str(r.values["registration_number"]): r.values for r in lido.rows}
    assert str(por_matricula[PESSOAS[0][0]]["performed_on"]).startswith("2025-09-02")


def test_valor_fora_do_catalogo_nao_derruba_a_linha_mas_aparece(
    ensaio: tuple[Resultado, Path],
) -> None:
    resultado, _ = ensaio
    campos = {(n.sheet, n.reason) for n in resultado.notas}

    assert ("DADOS FUNCIONÁRIOS", "REGIME") in campos
    assert ("VENCIMENTO ASO", "TIPO") in campos


def test_sem_data_de_reajuste_a_vigencia_e_a_admissao(ensaio: tuple[Resultado, Path]) -> None:
    resultado, saida = ensaio

    assert resultado.admissao_como_vigencia == 1
    template = get_template("hr_compensation")
    assert template is not None
    lido = parse(
        next(saida.glob("*-hr_compensation.xlsx")).read_bytes(),
        template=template,
        tenant_id=TENANT,
    )
    por_matricula = {str(r.values["registration_number"]): r.values for r in lido.rows}
    # A quinta pessoa da fixture é a que não tem data de reajuste.
    assert por_matricula[PESSOAS[4][0]]["effective_from"] is not None


def test_o_parqueado_agrupa_por_linha_de_origem_e_nao_por_pessoa(
    ensaio: tuple[Resultado, Path],
) -> None:
    """Três atestados da mesma pessoa são três registros, não um.

    Agrupar por matrícula parece uma limpeza e é uma perda: o último afastamento
    sobrescreveria os dois anteriores sem que ninguém visse.
    """
    _, saida = ensaio
    afastamentos = (
        (saida / "parqueado" / "leave_period.csv").read_text(encoding="utf-8").splitlines()[1:]
    )
    parcelas = (
        (saida / "parqueado" / "agreement_installment.csv")
        .read_text(encoding="utf-8")
        .splitlines()[1:]
    )

    assert sum(1 for linha in afastamentos if linha.startswith(PESSOAS[0][0])) == 5
    assert sum(1 for linha in parcelas if linha.startswith(PESSOAS[0][0])) == 3


# ---------------------------------------------------------------------------
# O mapa não pode divergir do sistema
# ---------------------------------------------------------------------------
def test_o_mapa_cobre_as_dezessete_abas_declaradas() -> None:
    esperadas = {
        "DADOS FUNCIONARIOS",
        "DESLIGADOS",
        "FALTAS_ATESTADOS",
        "MOVIMENTACOES",
        "FERIAS",
        "CNH",
        "VENCIMENTO ASO",
        "SINISTROS",
        "CONTROLE PARC",
        "UNIDADES",
        "COD POSTOS",
        "QUADRO GERAL",
        "GERAL",
        "FACE GERAL",
        "QUADRO_POSTOS",
        "CESTAS",
        "PLANILHA1",
    }

    assert {normalize(sheet.name) for sheet in SHEETS} == esperadas


def test_todo_destino_do_mapa_existe_no_sistema() -> None:
    """O mapa nomeia template e coluna: os dois têm de existir de verdade.

    É o mesmo princípio do `95_teste_matriz_rh.py`, um nível acima — lá as
    colunas são conferidas contra o banco, aqui os destinos são conferidos contra
    os templates e a matriz. Um mapa que aponta para coluna que não existe só
    falha na implantação, que é o pior lugar.
    """
    conhecidas = tables()

    for sheet in SHEETS:
        for column in sheet.columns:
            alvo = column.target
            if alvo == carga_inicial.CHAVE:
                continue
            if alvo is None:
                assert column.refused, f"{sheet.name}.{column.label}: sem destino e sem motivo"
                continue
            if alvo.startswith("@fallback:"):
                alvo = alvo.split(":", 1)[1]
            destino, coluna = alvo.split(".", 1)
            if destino == "conferencia":
                continue
            template = TEMPLATES.get(destino)
            if template is not None:
                assert coluna in {c.column for c in template.columns}, f"{alvo} fora do template"
            else:
                assert destino in conhecidas, f"{alvo}: {destino} não é tabela da matriz"


def test_toda_traducao_de_valor_cai_no_catalogo_do_banco() -> None:
    """A tradução para inglês tem de cair no catálogo que o banco aceita.

    A tradução é a única parte do conversor que inventa uma string. Um `periodico`
    no lugar de `periodic` passaria por aqui, atravessaria o parqueado inteiro e
    só falharia no dia em que a tabela ganhasse caminho de escrita — meses depois
    da implantação, com a origem já aposentada.
    """
    for sheet in SHEETS:
        for column in sheet.columns:
            if column.values is None or not column.target:
                continue
            destino, coluna = column.target.split(".", 1)
            template = TEMPLATES.get(destino)
            if template is not None:
                tabela = next(c.table for c in template.columns if c.column == coluna)
            else:
                tabela = destino
            catalogo = ENUMS.get((tabela, coluna))
            assert catalogo, f"{column.target}: traduz valor para coluna sem catálogo declarado"
            fora = set(column.values.values()) - catalogo
            assert not fora, f"{column.target}: {sorted(fora)} fora de ownership.ENUMS"


def test_aba_que_nao_vira_template_diz_por_que() -> None:
    for sheet in SHEETS:
        if sheet.purpose is not Purpose.TEMPLATE:
            assert sheet.reason, f"{sheet.name}: sem motivo declarado"


def test_o_conversor_nao_fala_com_o_banco() -> None:
    """Guarda sintática, no espírito do `bind_tenant`.

    O script roda na implantação, com a planilha real na mesa. Ele não pode ter
    credencial nem conexão — e a forma de garantir isso não é lembrar, é falhar.
    """
    fonte = inspect.getsource(carga_inicial)

    for proibido in ("psycopg", "operax.core.db", "service_role", "DATABASE_URL"):
        assert proibido not in fonte


# ---------------------------------------------------------------------------
# Recusas da corrida inteira
# ---------------------------------------------------------------------------
def test_diretorio_sem_modelo_para_a_corrida(tmp_path: Path) -> None:
    planilha = escrever_planilha(tmp_path / "cliente.xlsx")
    (tmp_path / "vazio").mkdir()

    with pytest.raises(ConverterError, match="nenhum modelo"):
        converter(planilha=planilha, modelos=tmp_path / "vazio", saida=tmp_path / "saida")


def test_arquivo_que_nao_e_modelo_do_sistema_e_recusado(tmp_path: Path) -> None:
    from openpyxl import Workbook

    planilha = escrever_planilha(tmp_path / "cliente.xlsx")
    modelos = escrever_modelos(tmp_path / "modelos")
    intruso = Workbook()
    intruso.active["A1"] = "planilha qualquer"
    intruso.save(modelos / "aaa-nao-e-modelo.xlsx")

    with pytest.raises(Exception, match="aba de controle"):
        converter(planilha=planilha, modelos=modelos, saida=tmp_path / "saida")


def test_dois_modelos_do_mesmo_tipo_param_a_corrida(tmp_path: Path) -> None:
    planilha = escrever_planilha(tmp_path / "cliente.xlsx")
    modelos = escrever_modelos(tmp_path / "modelos")
    copia = modelos / "zz-copia.xlsx"
    copia.write_bytes((modelos / "modelo-hr_link.xlsx").read_bytes())

    with pytest.raises(ConverterError, match="dois modelos"):
        converter(planilha=planilha, modelos=modelos, saida=tmp_path / "saida")


def test_modelo_de_outro_tenant_nao_e_preenchido(tmp_path: Path) -> None:
    """O conversor preenche; quem recusa o tenant é o import, e ele recusa.

    Vale escrever: a implantação de um cliente com os modelos de outro na mesma
    máquina é um cenário real, e o desfecho não pode depender de ninguém reparar
    no nome do arquivo.
    """
    outro = UUID("dede0000-0000-0000-0000-0000000000ff")
    planilha = escrever_planilha(tmp_path / "cliente.xlsx")
    modelos = escrever_modelos(tmp_path / "modelos", tenant_id=outro)
    saida = tmp_path / "saida"
    converter(planilha=planilha, modelos=modelos, saida=saida)

    template = get_template("hr_link")
    assert template is not None
    with pytest.raises(Exception, match="outro cliente"):
        parse(
            next(saida.glob("*-hr_link.xlsx")).read_bytes(),
            template=template,
            tenant_id=TENANT,
        )


def test_a_linha_de_comando_devolve_zero_quando_o_balanco_fecha(tmp_path: Path) -> None:
    planilha = escrever_planilha(tmp_path / "cliente.xlsx")
    modelos = escrever_modelos(tmp_path / "modelos")
    saida = tmp_path / "saida"

    codigo = carga_inicial.main(
        ["--planilha", str(planilha), "--modelos", str(modelos), "--saida", str(saida)]
    )

    assert codigo == 0
    assert (saida / "RELATORIO-DE-DESCARTE.md").exists()
    assert sorted(p.name for p in saida.glob("*.xlsx")) == [
        "01-hr_link.xlsx",
        "02-hr_employee.xlsx",
        "03-hr_exam.xlsx",
        "04-hr_compensation.xlsx",
    ]


def test_a_linha_de_comando_recusa_caminho_que_nao_existe(tmp_path: Path) -> None:
    codigo = carga_inicial.main(
        [
            "--planilha",
            str(tmp_path / "nao-existe.xlsx"),
            "--modelos",
            str(tmp_path),
            "--saida",
            str(tmp_path / "saida"),
        ]
    )

    assert codigo == 2


def test_o_exame_nao_herda_metade_do_pre_preenchimento(tmp_path: Path) -> None:
    """A linha vai inteira: o que a planilha não disse é apagado, não herdado.

    O caso que produziu isto: o cliente tinha um periódico de setembro gravado, a
    planilha trazia um exame de mudança de função em fevereiro sem validade, e o
    modelo saiu com a data de fevereiro e a validade de setembro — um exame que
    nunca existiu, com a validade de outro. Numa coluna de vencimentos, é a data
    errada na tela de quem persegue prazo de ASO.
    """
    planilha = escrever_planilha(tmp_path / "cliente.xlsx")
    modelos = escrever_modelos(
        tmp_path / "modelos",
        gravado={
            matricula: {
                "type": "periodic",
                "performed_on": date(2025, 9, 10),
                "valid_until": date(2026, 9, 10),
                "result": "fit",
            }
            for matricula in (PESSOAS[3][0], PESSOAS[4][0])
        },
    )
    saida = tmp_path / "saida"
    converter(planilha=planilha, modelos=modelos, saida=saida)

    template = get_template("hr_exam")
    assert template is not None
    lido = parse(
        next(saida.glob("*-hr_exam.xlsx")).read_bytes(), template=template, tenant_id=TENANT
    )
    por_matricula = {str(r.values["registration_number"]): r.values for r in lido.rows}

    # A planilha traz tipo, data e aptidão para esta pessoa — e nenhuma validade.
    sem_validade = por_matricula[PESSOAS[3][0]]
    assert str(sem_validade["performed_on"]).startswith("2026-02-10")
    assert sem_validade["type"] == "job_change"
    assert sem_validade["result"] == "fit"
    assert sem_validade["valid_until"] is None

    # E aqui o tipo da planilha ("EXAME DE RETORNO") não está no catálogo: a
    # célula fica vazia e o preview recusa por campo obrigatório. Herdar o
    # `periodic` que estava gravado seria inventar o tipo do exame.
    sem_tipo = por_matricula[PESSOAS[4][0]]
    assert sem_tipo["type"] is None
    assert str(sem_tipo["performed_on"]).startswith("2026-06-01")
