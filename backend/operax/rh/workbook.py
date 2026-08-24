"""The spreadsheet, written and read back — the only place that speaks .xlsx.

The file the customer downloads and the file they upload are the same object seen
twice, so building and parsing live together: a change to the header on one side
that the other side does not know about is impossible to write here without
seeing it.

WHAT THE PROTECTION IS AND IS NOT
Locking the sheet is a nudge. Excel's sheet protection has no password here and
anyone can turn it off; that is fine, because it is not the boundary. The
boundary is `check_owned_fields` on the server, which refuses the line whatever
the spreadsheet allowed. The lock exists so that nobody edits a Secullum column
by accident and only finds out at upload.

WHAT `_meta` IS FOR
An .xlsx has no identity. Without the hidden sheet, a file that lost a column, or
came from another tenant, or was generated three layouts ago, is indistinguishable
from a good one — and "importar no melhor esforço" is how a spreadsheet writes the
wrong person's salary. So the sheet carries the layout version, the type, the
tenant, and a hash of the header as generated. Any of them off, and the answer is
"baixe o modelo atual" before a single line is read.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from typing import Any
from uuid import UUID

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from operax.rh.templates import Template, accepted_versions

META_SHEET = "_meta"
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_BAIXE_O_MODELO = "Baixe o modelo atual e refaça o preenchimento."

# Cinza para a coluna que o Secullum governa, laranja de marca no cabeçalho.
# `FF8C00` é preenchimento com texto escuro por cima, nunca o contrário.
_FILL_HEADER = PatternFill("solid", fgColor="FF8C00")
_FILL_LOCKED = PatternFill("solid", fgColor="EFEFEF")
_FONT_HEADER = Font(bold=True, color="262626")
_UNLOCKED = Protection(locked=False)

_MIN_WIDTH = 12
_MAX_WIDTH = 42


class WorkbookError(RuntimeError):
    """O arquivo inteiro é recusado. Não é erro de linha — não há linha a ler."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True, slots=True)
class SheetRow:
    """Uma linha da aba de dados, com o número que o usuário vê no Excel."""

    line: int
    values: dict[str, Any]


@dataclass(frozen=True, slots=True)
class ParsedWorkbook:
    layout_version: str
    rows: tuple[SheetRow, ...]


def header_hash(labels: tuple[str, ...]) -> str:
    """Impressão digital do cabeçalho como foi gerado.

    Cobre rótulo, ordem e quantidade — que é exatamente o conjunto de coisas que
    muda quando alguém insere, remove ou renomeia uma coluna. Editar célula de
    dado não mexe nisto, e é o que o arquivo existe para permitir.
    """
    return hashlib.sha256("\n".join(labels).encode("utf-8")).hexdigest()


# ---------------------------------------------------------------------------
# Escrita
# ---------------------------------------------------------------------------
def build(
    template: Template,
    rows: list[dict[str, Any]],
    *,
    tenant_id: UUID,
    generated_at: datetime,
) -> bytes:
    """O modelo preenchido com o que já está gravado, pronto para editar."""
    wb = Workbook()
    ws = wb.active
    ws.title = template.sheet_title

    for indice, column in enumerate(template.columns, start=1):
        celula = ws.cell(row=1, column=indice, value=column.label)
        celula.font = _FONT_HEADER
        celula.fill = _FILL_HEADER
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if template.is_locked(column) and column.mirror:
            # A tela diz "Secullum · leitura de HH:MM"; aqui a mesma frase mora
            # no comentário da célula, que é onde o usuário do Excel procura.
            celula.comment = _comentario(f"origem: Secullum ({column.mirror}). Não editável.")
        elif template.is_locked(column):
            celula.comment = _comentario("Identifica a linha. Não editável.")
        largura = max(len(column.label) + 4, _MIN_WIDTH)
        ws.column_dimensions[get_column_letter(indice)].width = min(largura, _MAX_WIDTH)

    for deslocamento, registro in enumerate(rows):
        linha = 2 + deslocamento
        for indice, column in enumerate(template.columns, start=1):
            celula = ws.cell(row=linha, column=indice, value=registro.get(column.column))
            if template.is_locked(column):
                celula.fill = _FILL_LOCKED
            else:
                celula.protection = _UNLOCKED
            if column.kind == "date":
                celula.number_format = "DD/MM/YYYY"
            elif column.kind == "decimal":
                celula.number_format = "#,##0.00"

    ultima = max(2, 1 + len(rows))
    for indice, column in enumerate(template.columns, start=1):
        if template.is_locked(column) or not column.enum:
            continue
        letra = get_column_letter(indice)
        validacao = DataValidation(
            type="list",
            formula1='"{}"'.format(",".join(sorted(column.enum))),
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Valor fora do catálogo",
            error=f"{column.label} só aceita: {', '.join(sorted(column.enum))}",
        )
        ws.add_data_validation(validacao)
        validacao.add(f"{letra}2:{letra}{ultima}")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(template.columns))}{ultima}"
    ws.protection.sheet = True

    _escrever_meta(
        wb,
        template=template,
        tenant_id=tenant_id,
        generated_at=generated_at,
    )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _comentario(texto: str):
    from openpyxl.comments import Comment

    comentario = Comment(texto, "OperaX")
    comentario.width = 260
    comentario.height = 80
    return comentario


def _escrever_meta(
    wb: Workbook, *, template: Template, tenant_id: UUID, generated_at: datetime
) -> None:
    meta = wb.create_sheet(META_SHEET)
    pares = (
        ("layout_version", template.layout_version),
        ("type", template.type),
        ("tenant", str(tenant_id)),
        ("sheet", template.sheet_title),
        ("generated_at", generated_at.isoformat()),
        ("columns", ",".join(c.column for c in template.columns)),
        ("header_hash", header_hash(template.labels())),
    )
    for linha, (chave, valor) in enumerate(pares, start=1):
        meta.cell(row=linha, column=1, value=chave)
        meta.cell(row=linha, column=2, value=valor)
    meta.protection.sheet = True
    meta.sheet_state = "hidden"


# ---------------------------------------------------------------------------
# Leitura
# ---------------------------------------------------------------------------
def parse(data: bytes, *, template: Template, tenant_id: UUID) -> ParsedWorkbook:
    """O arquivo de volta, ou a recusa de arquivo inteiro com o motivo.

    A ordem das recusas é a da SPEC §3: o arquivo é julgado antes da primeira
    linha, porque um arquivo errado não tem linha certa.
    """
    try:
        wb = load_workbook(BytesIO(data), data_only=True, read_only=False)
    except Exception as erro:  # noqa: BLE001 — qualquer falha aqui é "não é xlsx"
        raise WorkbookError(
            "arquivo_ilegivel", f"Não consegui abrir o arquivo como .xlsx. {_BAIXE_O_MODELO}"
        ) from erro

    if META_SHEET not in wb.sheetnames:
        raise WorkbookError(
            "sem_meta",
            f"Arquivo sem a aba de controle `{META_SHEET}` — ele não foi gerado pelo sistema. "
            f"{_BAIXE_O_MODELO}",
        )

    meta = _ler_meta(wb[META_SHEET])

    if meta.get("type") != template.type:
        raise WorkbookError(
            "tipo_divergente",
            f"O arquivo é do tipo {meta.get('type')!r} e foi enviado como {template.type!r}. "
            "Escolha o tipo correto ou baixe o modelo do tipo desejado.",
        )
    if meta.get("tenant") != str(tenant_id):
        # Vale sozinho: um arquivo de outro cliente traz nome e matrícula de
        # outro cliente, e o preview mostraria essa gente na tela.
        raise WorkbookError(
            "tenant_divergente",
            "Este arquivo foi gerado para outro cliente e não pode ser importado aqui.",
        )
    if meta.get("layout_version") not in accepted_versions(template):
        raise WorkbookError(
            "layout_antigo",
            f"O modelo é da versão {meta.get('layout_version')!r}, que não é mais aceita. "
            f"{_BAIXE_O_MODELO}",
        )

    ws = wb[meta["sheet"]] if meta.get("sheet") in wb.sheetnames else _primeira_aba(wb)
    cabecalho = tuple(_texto(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1)))
    if header_hash(cabecalho) != meta.get("header_hash"):
        raise WorkbookError(
            "cabecalho_alterado",
            f"O cabeçalho da planilha foi alterado — coluna inserida, removida ou renomeada. "
            f"{_BAIXE_O_MODELO}",
        )

    chaves = meta.get("columns", "").split(",")
    if len(chaves) != len(cabecalho):
        raise WorkbookError(
            "cabecalho_alterado",
            f"O cabeçalho tem {len(cabecalho)} coluna(s) e o modelo declara {len(chaves)}. "
            f"{_BAIXE_O_MODELO}",
        )

    linhas: list[SheetRow] = []
    for numero, celulas in enumerate(ws.iter_rows(min_row=2), start=2):
        valores = {chave: celula.value for chave, celula in zip(chaves, celulas, strict=False)}
        if all(_texto(v) == "" for v in valores.values()):
            # Linha em branco no fim da planilha é o rastro de quem apagou o
            # conteúdo em vez de excluir a linha. Não é erro; não é linha.
            continue
        linhas.append(SheetRow(line=numero, values=valores))

    return ParsedWorkbook(layout_version=meta["layout_version"], rows=tuple(linhas))


def read_meta(wb: Workbook) -> dict[str, str]:
    """A identidade do arquivo, para quem precisa lê-la antes de escolher o template.

    `parse` já exige saber qual template esperar; o conversor de implantação não
    sabe — ele recebe um diretório de modelos baixados e descobre o que cada um é
    abrindo a aba de controle. Mesmo leitor, mesmas recusas.
    """
    if META_SHEET not in wb.sheetnames:
        raise WorkbookError(
            "sem_meta",
            f"Arquivo sem a aba de controle `{META_SHEET}` — ele não foi gerado pelo sistema. "
            f"{_BAIXE_O_MODELO}",
        )
    return _ler_meta(wb[META_SHEET])


def _primeira_aba(wb: Workbook):
    for nome in wb.sheetnames:
        if nome != META_SHEET:
            return wb[nome]
    raise WorkbookError("sem_dados", f"Arquivo sem aba de dados. {_BAIXE_O_MODELO}")


def _ler_meta(ws) -> dict[str, str]:
    meta: dict[str, str] = {}
    for linha in ws.iter_rows(min_row=1, max_col=2):
        chave = _texto(linha[0].value)
        if chave:
            meta[chave] = _texto(linha[1].value) if len(linha) > 1 else ""
    faltando = {"layout_version", "type", "tenant", "columns", "header_hash"} - set(meta)
    if faltando:
        raise WorkbookError(
            "meta_incompleta",
            f"A aba de controle está incompleta (falta {', '.join(sorted(faltando))}). "
            f"{_BAIXE_O_MODELO}",
        )
    return meta


def _texto(valor: Any) -> str:
    return "" if valor is None else str(valor).strip()
