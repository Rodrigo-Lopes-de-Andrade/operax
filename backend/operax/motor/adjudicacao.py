"""Human adjudication — the shadow gate's reference truth since 2026-09-09.

WHY THIS EXISTS
`docs/SPEC-TECNICA.md` §3.5 measured false positives against "what the Secullum
registered". The mirror never carries a verdict, only input, and going to the
source for one was measured and dropped: on 2.121 rows of `secullum."Batida"` in
production, `"Ajuste"` is filled once, `"Abono2".."Abono4"` and `"Observacoes"`
never, and the label we derive ourselves covers 4%. The customer does not
justify days in the Secullum, so its calculation is a deterministic function of
the same punches, rosters and tolerances this engine already reads — it would
agree with us almost everywhere, including on the supervisors who never punch,
where agreement *is* the false positive. See
`docs/DECISAO-VERDADE-DE-REFERENCIA-G4.md`.

WHY A SPREADSHEET AND NOT A SCREEN
The population is small enough to be a census, not a sample: 820 active shadow
events over 326 employee-days. A screen would be the right answer for a
recurring product task; this is a gate that runs twice. The round trip reuses
`openpyxl`, already a dependency, and adds no route, no RLS surface for the
browser and no frontend.

WHY THE SHEET SPEAKS PORTUGUESE AND THE TABLE SPEAKS THE VOCABULARY
Nobody types `exempt_from_punching` correctly at row 300. The sheet offers
pt-BR labels in a dropdown and this module translates; an unknown word is a
refused line, never a guessed one.

WHY AN INCOMPLETE FILE IS NOT AN ERROR, AND AN INVALID ONE REFUSES EVERYTHING
A blank verdict is a case nobody judged yet — normal halfway through a census,
and what `measure` reports as coverage. A *filled* line that contradicts itself
(false without cause, true with one) is a different thing: it means the person
and the vocabulary disagree, and importing the rest would bury that under a
number. So parsing collects every line error and writes nothing until they are
gone.

WHY COVERAGE TRAVELS WITH THE RATE, ALWAYS
A rate over a partial census is the false green this repository keeps finding:
5% of the 40 cases somebody bothered to judge says nothing about the 820. The
measurement refuses to call itself a gate result until every event in the window
has a verdict, and prints the two numbers side by side either way.
"""

from __future__ import annotations

import argparse
import unicodedata
from dataclasses import dataclass
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from operax.core.db import run_cli
from operax.core.tenant import SystemContext, active_tenants, tenant_scope

TASK = "adjudicacao"

#: The gate of regra 8, and the number the SPEC §3.5 step 5 names.
GATE_MAX_FALSE_POSITIVE = 5.0

#: Default window: the shadow stage runs on "1 to 2 weeks of real data" (§3.5).
DEFAULT_WINDOW_DAYS = 30

VERDICT_PT: dict[str, str] = {
    "verdadeiro": "true_positive",
    "falso": "false_positive",
}

CAUSE_PT: dict[str, str] = {
    "escala errada": "wrong_schedule",
    "tolerancia errada": "wrong_tolerance",
    "erro do motor": "engine_bug",
    "nao bate ponto": "exempt_from_punching",
    "justificado fora do sistema": "justified_outside_system",
}

HEADERS: tuple[str, ...] = (
    "indicio_id",
    "colaborador",
    "unidade",
    "dia",
    "tipo",
    "minutos",
    "esperado",
    "realizado",
    "batidas do dia",
    "jornada (min)",
    "confianca da escala",
    "veredito",
    "causa (so se falso)",
    "observacao",
)

_ID_COLUMN = 1
_VERDICT_COLUMN = 12
_CAUSE_COLUMN = 13
_NOTE_COLUMN = 14


def _fold(value: object) -> str:
    """Casefold and strip accents, so `Tolerância` and `tolerancia` are one word."""
    text = str(value or "").strip().casefold()
    return "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))


# ---------------------------------------------------------------------------
# The census
# ---------------------------------------------------------------------------

_CENSUS_SQL = """
with ponte as (
    -- The punches are keyed by the mirror's employee, never by ours. Same
    -- bridge `regras.py` uses, and it stays a left join: a person the mirror
    -- lost still gets a row here, with an empty punch list that is itself the
    -- evidence the judge needs.
    select e.id as employee_id, f.id as mirror_id
      from app.employee e
      left join secullum."Funcionario" f
             on f.tenant_id = e.tenant_id
            and f."FuncionarioId" = e.secullum_employee_id
     where e.tenant_id = %(tenant_id)s
)
select d.id                                as deviation_event_id,
       e.name                              as employee_name,
       coalesce(u.name, '-')               as unit_name,
       d.reference_date                    as reference_date,
       d.type                              as type,
       d.minutes                           as minutes,
       to_char(d.expected_time, 'HH24:MI') as expected_time,
       to_char(d.actual_time, 'HH24:MI')   as actual_time,
       coalesce((
           select string_agg(to_char(m.hora, 'HH24:MI'), ' ' order by m.hora)
             from app.batida_marcacao m
             join ponte p on p.mirror_id = m.funcionario_id
            where p.employee_id = d.employee_id
              and m.data = d.reference_date
              and m.hora is not null
              and not m.desconsiderada
       ), '')                              as punches,
       w.workload_minutes                  as workload_minutes,
       w.confidence                        as confidence,
       a.verdict                           as verdict,
       a.cause                             as cause,
       a.note                              as note
  from app.deviation_event d
  join app.employee e on e.id = d.employee_id
  left join app.unit u on u.id = d.unit_id
  left join app.expected_workday w
         on w.tenant_id = d.tenant_id
        and w.employee_id = d.employee_id
        and w.reference_date = d.reference_date
  left join app.deviation_adjudication a on a.deviation_event_id = d.id
 where d.tenant_id = %(tenant_id)s
   and d.mode = %(mode)s
   and d.status = 'active'
   and d.reference_date between %(start)s::date and %(end)s::date
 order by e.name, d.reference_date, d.type
"""


async def census(
    context: SystemContext, start: date, end: date, *, mode: str = "shadow"
) -> list[dict[str, Any]]:
    """Every active event of the window, with the day around it.

    The context columns are not decoration: "escala errada" and "tolerância
    errada" are two of the five causes, and a judge who cannot see the roster
    behind the event can only ever answer "erro do motor".
    """
    async with tenant_scope(context) as scope:
        await scope.execute(_CENSUS_SQL, {"start": start, "end": end, "mode": mode})
        return [dict(row) for row in await scope.fetchall()]


def build_workbook(rows: list[dict[str, Any]]) -> bytes:
    """The sheet the judge fills. Already-judged rows come back pre-filled."""
    wb = Workbook()
    ws = wb.active
    ws.title = "censo"

    for column, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    pt_por_verdict = {v: k for k, v in VERDICT_PT.items()}
    pt_por_cause = {v: k for k, v in CAUSE_PT.items()}

    for line, row in enumerate(rows, start=2):
        ws.cell(row=line, column=1, value=str(row["deviation_event_id"]))
        ws.cell(row=line, column=2, value=row["employee_name"])
        ws.cell(row=line, column=3, value=row["unit_name"])
        ws.cell(row=line, column=4, value=row["reference_date"].strftime("%d/%m/%Y"))
        ws.cell(row=line, column=5, value=row["type"])
        ws.cell(row=line, column=6, value=row["minutes"])
        ws.cell(row=line, column=7, value=row["expected_time"] or "")
        ws.cell(row=line, column=8, value=row["actual_time"] or "")
        ws.cell(row=line, column=9, value=row["punches"])
        ws.cell(row=line, column=10, value=row["workload_minutes"])
        ws.cell(row=line, column=11, value=row["confidence"])
        ws.cell(row=line, column=12, value=pt_por_verdict.get(row["verdict"] or "", ""))
        ws.cell(row=line, column=13, value=pt_por_cause.get(row["cause"] or "", ""))
        ws.cell(row=line, column=14, value=row["note"] or "")

    ultima = max(len(rows) + 1, 2)
    verdict_dv = DataValidation(
        type="list", formula1='"' + ",".join(VERDICT_PT) + '"', allow_blank=True
    )
    cause_dv = DataValidation(
        type="list", formula1='"' + ",".join(CAUSE_PT) + '"', allow_blank=True
    )
    ws.add_data_validation(verdict_dv)
    ws.add_data_validation(cause_dv)
    verdict_dv.add(f"L2:L{ultima}")
    cause_dv.add(f"M2:M{ultima}")

    larguras = (38, 28, 18, 12, 20, 9, 10, 10, 30, 13, 18, 14, 26, 40)
    for column, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(column)].width = largura
    ws.freeze_panes = "B2"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Reading it back
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class ParsedLine:
    line: int
    deviation_event_id: UUID
    verdict: str
    cause: str | None
    note: str | None


@dataclass(frozen=True)
class LineError:
    line: int
    message: str


def parse(data: bytes) -> tuple[list[ParsedLine], list[LineError]]:
    """Per-line verdict on the file itself. Writes nothing, ever."""
    wb = load_workbook(BytesIO(data), data_only=True)
    ws = wb[wb.sheetnames[0]]

    cabecalho = [_fold(ws.cell(row=1, column=c).value) for c in range(1, len(HEADERS) + 1)]
    if cabecalho[_ID_COLUMN - 1] != _fold(HEADERS[0]) or cabecalho[_VERDICT_COLUMN - 1] != _fold(
        HEADERS[_VERDICT_COLUMN - 1]
    ):
        return [], [LineError(1, "cabeçalho não confere: use a planilha gerada por `exportar`")]

    linhas: list[ParsedLine] = []
    erros: list[LineError] = []

    for line in range(2, ws.max_row + 1):
        bruto_id = str(ws.cell(row=line, column=_ID_COLUMN).value or "").strip()
        veredito = _fold(ws.cell(row=line, column=_VERDICT_COLUMN).value)
        causa = _fold(ws.cell(row=line, column=_CAUSE_COLUMN).value)
        nota = str(ws.cell(row=line, column=_NOTE_COLUMN).value or "").strip() or None

        if not bruto_id and not veredito:
            continue
        if not veredito:
            # Not judged yet. Legitimate, and `measure` counts it as coverage
            # missing rather than pretending it is a true positive.
            continue
        try:
            evento = UUID(bruto_id)
        except ValueError:
            erros.append(LineError(line, f"indicio_id inválido: {bruto_id!r}"))
            continue
        if veredito not in VERDICT_PT:
            erros.append(LineError(line, f"veredito {veredito!r} não é 'verdadeiro' nem 'falso'"))
            continue

        verdict = VERDICT_PT[veredito]
        if causa and causa not in CAUSE_PT:
            erros.append(LineError(line, f"causa {causa!r} está fora do vocabulário"))
            continue
        cause = CAUSE_PT[causa] if causa else None
        if verdict == "false_positive" and cause is None:
            erros.append(LineError(line, "falso positivo sem causa — o conserto fica sem endereço"))
            continue
        if verdict == "true_positive" and cause is not None:
            erros.append(LineError(line, "verdadeiro positivo não leva causa"))
            continue

        linhas.append(ParsedLine(line, evento, verdict, cause, nota))

    return linhas, erros


_SAVE_SQL = """
insert into app.deviation_adjudication
    (tenant_id, deviation_event_id, verdict, cause, note, author_name)
select d.tenant_id, d.id, %(verdict)s, %(cause)s, %(note)s, %(author_name)s
  from app.deviation_event d
 where d.id = %(deviation_event_id)s
   and d.tenant_id = %(tenant_id)s
on conflict (deviation_event_id) do update
   set verdict        = excluded.verdict,
       cause          = excluded.cause,
       note           = excluded.note,
       author_name    = excluded.author_name,
       adjudicated_at = now()
returning id
"""


async def save(context: SystemContext, lines: list[ParsedLine], *, author: str) -> list[ParsedLine]:
    """Write the verdicts of one tenant. Returns the lines it did NOT recognise.

    The insert reads `tenant_id` from the event instead of trusting the caller,
    and filters the event by the bound tenant: a row carrying an id from another
    customer writes nothing and comes back in the refused list, rather than
    landing a verdict of ours on top of somebody else's event.
    """
    desconhecidas: list[ParsedLine] = []
    async with tenant_scope(context) as scope:
        for line in lines:
            await scope.execute(
                _SAVE_SQL,
                {
                    "deviation_event_id": str(line.deviation_event_id),
                    "verdict": line.verdict,
                    "cause": line.cause,
                    "note": line.note,
                    "author_name": author,
                },
            )
            if await scope.fetchone() is None:
                desconhecidas.append(line)
    return desconhecidas


# ---------------------------------------------------------------------------
# The measurement
# ---------------------------------------------------------------------------

_MEASURE_SQL = """
select coalesce(d.type, '__total__')                            as type,
       count(*)                                                 as events,
       count(a.id)                                              as judged,
       count(*) filter (where a.verdict = 'false_positive')     as false_positives
  from app.deviation_event d
  left join app.deviation_adjudication a on a.deviation_event_id = d.id
 where d.tenant_id = %(tenant_id)s
   and d.mode = %(mode)s
   and d.status = 'active'
   and d.reference_date between %(start)s::date and %(end)s::date
 group by rollup (d.type)
 order by (d.type is null), d.type
"""


@dataclass(frozen=True)
class Slice:
    type: str
    events: int
    judged: int
    false_positives: int

    @property
    def complete(self) -> bool:
        return self.events > 0 and self.judged == self.events

    @property
    def rate(self) -> float | None:
        """False positives over what was judged — `None` while nothing was."""
        return 100.0 * self.false_positives / self.judged if self.judged else None


@dataclass(frozen=True)
class Measurement:
    tenant_id: UUID
    start: date
    end: date
    mode: str
    total: Slice
    by_type: list[Slice]

    @property
    def gate_passes(self) -> bool | None:
        """`None` means the gate has no answer yet — not that it failed.

        A partial census cannot pass and cannot fail: that is the whole reason
        coverage is carried next to the rate instead of behind it.
        """
        if not self.total.complete:
            return None
        return (self.total.rate or 0.0) <= GATE_MAX_FALSE_POSITIVE


async def measure(
    context: SystemContext, start: date, end: date, *, mode: str = "shadow"
) -> Measurement:
    async with tenant_scope(context) as scope:
        await scope.execute(_MEASURE_SQL, {"start": start, "end": end, "mode": mode})
        rows = [dict(row) for row in await scope.fetchall()]

    fatias = [
        Slice(row["type"], row["events"], row["judged"], row["false_positives"]) for row in rows
    ]
    total = next((f for f in fatias if f.type == "__total__"), Slice("__total__", 0, 0, 0))
    return Measurement(
        tenant_id=context.tenant_id,
        start=start,
        end=end,
        mode=mode,
        total=total,
        by_type=[f for f in fatias if f.type != "__total__"],
    )


def relatorio(medicoes: list[Measurement]) -> str:
    """O relatório do censo, em pt-BR porque quem lê é o operador."""
    linhas: list[str] = []
    for m in medicoes:
        linhas.append(
            f"tenant {m.tenant_id} · {m.start} a {m.end} · modo {m.mode}\n"
            f"  {m.total.judged}/{m.total.events} indício(s) julgado(s) "
            f"· {m.total.false_positives} falso(s) positivo(s)"
        )
        for f in m.by_type:
            taxa = f"{f.rate:.1f}%" if f.rate is not None else "—"
            linhas.append(
                f"    {f.type:<20} {f.judged:>4}/{f.events:<4} julgado(s) · "
                f"{f.false_positives:>3} falso(s) · {taxa:>6}"
            )
        if m.total.events == 0:
            linhas.append("  nenhum indício na janela — nada a julgar")
        elif m.gate_passes is None:
            faltam = m.total.events - m.total.judged
            parcial = f"{m.total.rate:.1f}%" if m.total.rate is not None else "—"
            linhas.append(
                f"  ⚠️  CENSO INCOMPLETO: faltam {faltam} veredito(s). A taxa parcial é "
                f"{parcial} e ela NÃO responde ao G4 — o gate pede a janela inteira."
            )
        else:
            taxa = m.total.rate or 0.0
            situacao = "PASSA" if m.gate_passes else "NÃO PASSA"
            linhas.append(
                f"  falso positivo {taxa:.1f}% sobre o censo completo · "
                f"gate ≤{GATE_MAX_FALSE_POSITIVE:.0f}%: {situacao}"
            )
            linhas.append(
                "  ⚠️  uma execução só não promove: a §3.5 pede duas seguidas abaixo do teto."
            )
    return "\n".join(linhas) or "nenhum tenant ativo"


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


# ⛔ A CLI É SÍNCRONA E O BANCO É ASSÍNCRONO, E A FRONTEIRA FICA AQUI DE PROPÓSITO
# Ler e escrever arquivo dentro de corrotina bloqueia o loop — o `ruff` recusa
# (ASYNC230) e ele está certo. Então cada comando faz o I/O de arquivo no
# corpo síncrono e chama `run_cli` só para o trecho que fala com o banco.


def _exportar(args: argparse.Namespace) -> int:
    end = date.today()
    start = end - timedelta(days=args.dias - 1)

    async def coletar() -> list[tuple[UUID, list[dict[str, Any]]]]:
        return [
            (context.tenant_id, await census(context, start, end, mode=args.modo))
            for context in await active_tenants(TASK)
        ]

    coletado = run_cli(coletar())
    for tenant_id, rows in coletado:
        destino = args.saida
        if len(coletado) > 1:
            destino = f"{args.saida.removesuffix('.xlsx')}-{tenant_id}.xlsx"
        Path(destino).write_bytes(build_workbook(rows))
        print(f"{destino}: {len(rows)} indício(s) de {start} a {end} (modo {args.modo})")
    return 0


def _importar(args: argparse.Namespace) -> int:
    linhas, erros = parse(Path(args.arquivo).read_bytes())

    if erros:
        print(f"{len(erros)} linha(s) recusada(s) — nada foi gravado:")
        for erro in erros:
            print(f"  linha {erro.line}: {erro.message}")
        return 1

    async def gravar() -> list[ParsedLine]:
        pendentes = list(linhas)
        for context in await active_tenants(TASK):
            if not pendentes:
                break
            pendentes = await save(context, pendentes, author=args.autor)
        return pendentes

    pendentes = run_cli(gravar())
    print(f"{len(linhas) - len(pendentes)} veredito(s) gravado(s) por {args.autor}.")
    if pendentes:
        print(f"⚠️  {len(pendentes)} linha(s) com indício que nenhum tenant reconhece:")
        for linha in pendentes:
            print(f"  linha {linha.line}: {linha.deviation_event_id}")
        return 1
    return 0


def _medir(args: argparse.Namespace) -> int:
    end = date.today()
    start = end - timedelta(days=args.dias - 1)

    async def medir_todos() -> list[Measurement]:
        return [
            await measure(context, start, end, mode=args.modo)
            for context in await active_tenants(TASK)
        ]

    print(relatorio(run_cli(medir_todos())))
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Censo de adjudicação do modo sombra — a verdade de referência do G4."
    )
    parser.add_argument("--modo", default="shadow", help="shadow (padrão) ou production")
    parser.add_argument(
        "--dias",
        type=int,
        default=DEFAULT_WINDOW_DAYS,
        help=f"janela (padrão {DEFAULT_WINDOW_DAYS})",
    )
    sub = parser.add_subparsers(dest="comando", required=True)

    exportar = sub.add_parser("exportar", help="gera a planilha do censo")
    exportar.add_argument("--saida", default="censo-adjudicacao.xlsx")
    exportar.set_defaults(func=_exportar)

    importar = sub.add_parser("importar", help="lê a planilha preenchida")
    importar.add_argument("--arquivo", required=True)
    importar.add_argument("--autor", required=True, help="quem julgou — vai para author_name")
    importar.set_defaults(func=_importar)

    medir = sub.add_parser("medir", help="taxa de falso positivo, com a cobertura ao lado")
    medir.set_defaults(func=_medir)

    args = parser.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
